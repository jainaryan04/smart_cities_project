"""
OpenXML styling utilities and helper functions for publication-grade Word documents.
Provides native Microsoft Word OMML math conversion (via Office's MML2OMML.XSL),
booktabs table styling, fixed grid layout enforcement, callout boxes, and algorithms.
"""

import os
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import lxml.etree as ET
from latex2mathml.converter import convert as latex2mathml_convert

# ---------------------------------------------------------
# CONSTANTS & PALETTE
# ---------------------------------------------------------
COLOR_PRIMARY_NAVY = RGBColor(0x1A, 0x36, 0x5D)    # #1A365D - Main Title & Section Headings
COLOR_SECONDARY_SLATE = RGBColor(0x2B, 0x4C, 0x7E) # #2B4C7E - Subsections
COLOR_BODY_TEXT = RGBColor(0x1A, 0x1A, 0x1A)       # #1A1A1A - High-contrast readable body
COLOR_MUTED_GRAY = RGBColor(0x55, 0x55, 0x55)      # #555555 - Captions, metadata, notes

HEX_NAVY = "1A365D"
HEX_LIGHT_HEADER_BG = "EBF3FA"
HEX_ALT_ROW_BG = "F8FAFC"
HEX_CALLOUT_BG = "F4F6F9"
HEX_BORDER_LIGHT = "D0D7DE"
HEX_DIVIDER = "CBD5E1"

# ---------------------------------------------------------
# XSLT INITIALIZATION FOR NATIVE WORD OMML MATH
# ---------------------------------------------------------
XSLT_PATH = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
if os.path.exists(XSLT_PATH):
    xslt_doc = ET.parse(XSLT_PATH)
    transform = ET.XSLT(xslt_doc)
else:
    transform = None
    print("WARNING: MML2OMML.XSL not found. Math will fallback to formatted text.")

# ---------------------------------------------------------
# XML STYLING HELPERS
# ---------------------------------------------------------
def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    """Set inner cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set cell background fill color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_booktabs_borders(table, border_color=HEX_NAVY):
    """Apply professional booktabs borders: top/bottom solid rules, light header rule."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER_LIGHT}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_borderless_table(table):
    """Remove all borders from a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_table_fixed_grid(table, col_widths_in_inches):
    """
    Enforces strict fixed table layout and explicit grid column widths in dxa.
    Guarantees column width stability across Word versions, LibreOffice, and Google Docs.
    """
    table.autofit = False
    tblPr = table._tbl.tblPr
    
    # 1. Fixed layout
    if tblPr.find(qn('w:tblLayout')) is None:
        tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        
    # 2. Total table width in dxa
    total_dxa = sum(int(round(w * 1440)) for w in col_widths_in_inches)
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is not None:
        tblW.set(qn('w:type'), 'dxa')
        tblW.set(qn('w:w'), str(total_dxa))
    else:
        tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_dxa}" w:type="dxa"/>'))
        
    # 3. Explicit w:tblGrid
    tblGrid = table._tbl.tblGrid
    if tblGrid is not None:
        while len(tblGrid.gridCol_lst) < len(col_widths_in_inches):
            tblGrid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="1440"/>'))
        for c_idx, w in enumerate(col_widths_in_inches):
            tblGrid.gridCol_lst[c_idx].w = int(round(w * 1440))
            
    # 4. Apply widths and cantSplit across all rows and cells
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
        for c_idx, cell in enumerate(row.cells):
            if c_idx < len(col_widths_in_inches):
                w_in = col_widths_in_inches[c_idx]
                cell.width = Inches(w_in)
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is not None:
                    tcW.set(qn('w:type'), 'dxa')
                    tcW.set(qn('w:w'), str(int(round(w_in * 1440))))
                else:
                    tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(round(w_in * 1440))}" w:type="dxa"/>'))

def make_callout_box(cell):
    """Style a 1-cell table as an executive callout box."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    set_cell_shading(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

def set_algorithm_box_borders(table):
    """Style an algorithm box with double rules top and bottom."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="double" w:sz="12" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:bottom w:val="double" w:sz="12" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

# ---------------------------------------------------------
# TEXT RUN & MATH PARSER
# ---------------------------------------------------------
def parse_markdown_runs(text):
    """
    Scans a string and returns a sequence of tokens:
    ('text', content, is_bold, is_italic)
    ('math', latex_str, is_bold, is_italic)
    """
    tokens = []
    i = 0
    n = len(text)
    cur_text = []
    bold = False
    italic = False
    
    while i < n:
        if text[i:i+2] == '**':
            if cur_text:
                tokens.append(('text', ''.join(cur_text), bold, italic))
                cur_text = []
            bold = not bold
            i += 2
        elif text[i] == '*' and (i + 1 == n or text[i+1] != '*'):
            if cur_text:
                tokens.append(('text', ''.join(cur_text), bold, italic))
                cur_text = []
            italic = not italic
            i += 1
        elif text[i] == '$' and (i + 1 < n and text[i+1] != '$'):
            if cur_text:
                tokens.append(('text', ''.join(cur_text), bold, italic))
                cur_text = []
            j = text.find('$', i + 1)
            if j != -1:
                math_content = text[i+1:j]
                tokens.append(('math', math_content, bold, italic))
                i = j + 1
            else:
                cur_text.append(text[i])
                i += 1
        else:
            cur_text.append(text[i])
            i += 1
            
    if cur_text:
        tokens.append(('text', ''.join(cur_text), bold, italic))
        
    return tokens

def add_formatted_runs_to_paragraph(p, text, font_size=Pt(10.5), font_name='Times New Roman',
                                    default_color=COLOR_BODY_TEXT, force_bold=False, force_italic=False):
    """Appends formatted runs (bold, italic, inline math) to a python-docx paragraph."""
    tokens = parse_markdown_runs(text)
    for kind, content, is_bold, is_italic in tokens:
        b = force_bold or is_bold
        it = force_italic or is_italic
        
        if kind == 'text':
            r = p.add_run(content)
            r.font.name = font_name
            r.font.size = font_size
            r.font.bold = b
            r.font.italic = it
            r.font.color.rgb = default_color
        elif kind == 'math':
            omml_elem = None
            if transform is not None:
                try:
                    mml = latex2mathml_convert(content)
                    omml_tree = transform(ET.fromstring(mml))
                    omml_elem = omml_tree.getroot()
                except Exception:
                    omml_elem = None
                    
            if omml_elem is not None:
                p._p.append(omml_elem)
            else:
                clean = (content.replace('\\pm', '±')
                                .replace('\\le', '≤')
                                .replace('\\ge', '≥')
                                .replace('\\times', '×')
                                .replace('\\sigma', 'σ')
                                .replace('\\epsilon', 'ε')
                                .replace('\\delta', 'δ')
                                .replace('\\mu', 'μ')
                                .replace('\\lambda', 'λ')
                                .replace('\\eta', 'η')
                                .replace('\\theta', 'θ')
                                .replace('\\dagger', '†')
                                .replace('\\ddagger', '‡')
                                .replace('\\mathrm', '')
                                .replace('\\overline', '')
                                .replace('{', '')
                                .replace('}', ''))
                r = p.add_run(clean)
                r.font.name = font_name
                r.font.size = font_size
                r.font.bold = b
                r.font.italic = True
                r.font.color.rgb = default_color

def format_cell_runs(p, text, font_size=Pt(7.0), default_color=COLOR_BODY_TEXT, force_bold=False):
    """
    Renders table cell text with high-fidelity publication typography:
    Clean Unicode math symbols at exact cell font size, with proper bold, italic, and superscripts.
    Eliminates awkward 11 pt math font jumping in small table cells.
    """
    t = text.strip()
    
    # Check if whole cell should be bold
    cell_bold = force_bold or (t.startswith('**') and t.endswith('**'))
    if t.startswith('**') and t.endswith('**') and len(t) > 4:
        t = t[2:-2].strip()
        
    # Standardize LaTeX math expressions to clean unicode representation
    t = t.replace('$^{\\dagger}$', '§SUP_DAGGER§')
    t = t.replace('$^{\\ddagger}$', '§SUP_DOUBLE_DAGGER§')
    t = t.replace('\\dagger', '§SUP_DAGGER§')
    t = t.replace('\\ddagger', '§SUP_DOUBLE_DAGGER§')
    t = t.replace('\\overline{\\mathrm{JSD}}', 'JSD̄')
    t = t.replace('\\overline{JSD}', 'JSD̄')
    t = t.replace('\\mathrm{JSD}', 'JSD')
    t = t.replace('\\pm', '±')
    t = t.replace('\\le', '≤')
    t = t.replace('\\ge', '≥')
    t = t.replace('\\times', '×')
    t = t.replace('\\sigma', 'σ')
    t = t.replace('\\epsilon', 'ε')
    t = t.replace('\\delta', 'δ')
    t = t.replace('\\mu', 'μ')
    t = t.replace('\\lambda', 'λ')
    t = t.replace('\\eta', 'η')
    t = t.replace('\\infty', '∞')
    t = t.replace('$-$', '−')
    t = t.replace('--', '–') # en-dash for ranges
    
    # Handle $p$-value
    t = t.replace('$p$-value', '§ITALIC_P§-value')
    
    # Strip remaining math delimiters
    t = re.sub(r'\$([^\$]+?)\$', r'\1', t)
    t = t.replace('{', '').replace('}', '')
    
    # Parse tokens for italics and superscripts
    pattern = r'(§ITALIC_P§|§SUP_DAGGER§|§SUP_DOUBLE_DAGGER§|\*[^*]+?\*)'
    parts = re.split(pattern, t)
    
    for part in parts:
        if not part:
            continue
        if part == '§ITALIC_P§':
            r = p.add_run('p')
            r.font.name = 'Times New Roman'
            r.font.size = font_size
            r.font.bold = cell_bold
            r.font.italic = True
            r.font.color.rgb = default_color
        elif part == '§SUP_DAGGER§':
            r = p.add_run('†')
            r.font.name = 'Times New Roman'
            r.font.size = font_size
            r.font.bold = cell_bold
            r.font.superscript = True
            r.font.color.rgb = default_color
        elif part == '§SUP_DOUBLE_DAGGER§':
            r = p.add_run('‡')
            r.font.name = 'Times New Roman'
            r.font.size = font_size
            r.font.bold = cell_bold
            r.font.superscript = True
            r.font.color.rgb = default_color
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.font.name = 'Times New Roman'
            r.font.size = font_size
            r.font.bold = cell_bold
            r.font.italic = True
            r.font.color.rgb = default_color
        else:
            r = p.add_run(part)
            r.font.name = 'Times New Roman'
            r.font.size = font_size
            r.font.bold = cell_bold
            r.font.color.rgb = default_color

def add_equation_block(doc, latex_str, eq_num_str):
    """
    Inserts a display equation centered with a right-aligned equation number.
    Uses a 1-row, 2-column borderless table with strict fixed grid (6.3 in + 0.7 in = 7.0 in)
    and native Office Math Markup Language (OMML) display formatting.
    """
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borderless_table(tbl)
    set_table_fixed_grid(tbl, [6.3, 0.7])
    
    row = tbl.rows[0]
    cell_eq = row.cells[0]
    cell_num = row.cells[1]
    
    set_cell_margins(cell_eq, top=40, bottom=40, left=20, right=20)
    set_cell_margins(cell_num, top=40, bottom=40, left=20, right=20)
    
    p_eq = cell_eq.paragraphs[0]
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq.paragraph_format.space_before = Pt(0)
    p_eq.paragraph_format.space_after = Pt(0)
    
    omml_elem = None
    if transform is not None:
        try:
            mml = latex2mathml_convert(latex_str, display='block')
            omml_tree = transform(ET.fromstring(mml))
            root = omml_tree.getroot()
            if root.tag.endswith('oMathPara'):
                child_math = root.find('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
                omml_elem = child_math if child_math is not None else root
            else:
                omml_elem = root
        except Exception as e:
            print(f"OMML error for eq {eq_num_str}: {e}")
            omml_elem = None
            
    if omml_elem is not None:
        p_eq._p.append(omml_elem)
    else:
        r = p_eq.add_run(latex_str)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.font.italic = True
        
    p_num = cell_num.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_before = Pt(0)
    p_num.paragraph_format.space_after = Pt(0)
    r_num = p_num.add_run(eq_num_str)
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(10)
    r_num.font.bold = False
    r_num.font.color.rgb = COLOR_BODY_TEXT
    
    # Subtle trailing spacing
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(3)

def add_custom_heading(doc, text, level):
    """Adds a beautifully styled academic heading."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY_NAVY
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY_SLATE
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_body_paragraph(doc, text, space_after=4, bullet_type=None, indent_level=0):
    """Adds a standard justified body paragraph with rich runs."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    
    if indent_level > 0:
        p.paragraph_format.left_indent = Inches(0.25 * indent_level)
        
    add_formatted_runs_to_paragraph(p, text, font_size=Pt(10.5))
    return p

def add_figure_block(doc, image_path, caption_title, caption_text, width=Inches(6.8)):
    """Embeds a high-resolution figure with caption centered beneath."""
    if not os.path.exists(image_path):
        print(f"WARNING: Image not found at {image_path}")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    
    r_img = p_img.add_run()
    r_img.add_picture(image_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.left_indent = Inches(0.2)
    p_cap.paragraph_format.right_indent = Inches(0.2)
    
    # Title in bold
    r_title = p_cap.add_run(caption_title + " ")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY_NAVY
    
    # Caption body in regular / italic
    add_formatted_runs_to_paragraph(p_cap, caption_text, font_size=Pt(9.0),
                                    default_color=COLOR_MUTED_GRAY, force_italic=True)

def add_academic_table(doc, headers, data_rows, caption_text, notes_text=None,
                       col_widths=None, font_size=Pt(8.0), group_interval=None):
    """
    Creates a publication-grade table with booktabs borders, alternating row fills,
    explicit fixed grid column widths, header repetition across page breaks, and notes.
    """
    # Caption above table
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.keep_with_next = True
    
    parts = caption_text.split(':', 1)
    if len(parts) == 2:
        r_lbl = p_cap.add_run(parts[0] + ": ")
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_PRIMARY_NAVY
        
        r_title = p_cap.add_run(parts[1].strip())
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(9.5)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_BODY_TEXT
    else:
        r_lbl = p_cap.add_run(caption_text)
        r_lbl.font.name = 'Times New Roman'
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_PRIMARY_NAVY

    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_booktabs_borders(table)
    
    # Apply fixed grid column widths
    if col_widths:
        set_table_fixed_grid(table, col_widths)
    
    # Format Header Row
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for c_idx, (cell, h_text) in enumerate(zip(hdr_row.cells, headers)):
        set_cell_shading(cell, HEX_LIGHT_HEADER_BG)
        set_cell_margins(cell, top=60, bottom=60, left=50, right=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        format_cell_runs(p, h_text, font_size=font_size,
                         default_color=COLOR_PRIMARY_NAVY, force_bold=True)
                         
    # Format Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        # Determine background shading
        if group_interval is not None and group_interval > 0:
            bg = "FFFFFF" if (r_idx // group_interval) % 2 == 0 else HEX_ALT_ROW_BG
        else:
            bg = "FFFFFF" if r_idx % 2 == 0 else HEX_ALT_ROW_BG
            
        for c_idx, (cell, val_text) in enumerate(zip(row.cells, row_data)):
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=40, bottom=40, left=45, right=45)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            # Highlight bold rows (like Global Pooled)
            is_bold = '**' in val_text or 'Global Pooled' in val_text or 'This Study' in val_text
            format_cell_runs(p, val_text, font_size=font_size,
                             default_color=COLOR_BODY_TEXT, force_bold=is_bold)
                             
    # Notes below table
    if notes_text:
        p_notes = doc.add_paragraph()
        p_notes.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_notes.paragraph_format.space_before = Pt(3)
        p_notes.paragraph_format.space_after = Pt(10)
        add_formatted_runs_to_paragraph(p_notes, notes_text, font_size=Pt(8.0),
                                        default_color=COLOR_MUTED_GRAY, force_italic=True)

def add_algorithm_block(doc, title, require_text, ensure_text, steps):
    """Renders a beautifully boxed publication-grade algorithm block."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_algorithm_box_borders(tbl)
    set_table_fixed_grid(tbl, [7.0])
    
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, "FAFAFA")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    p_title = cell.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY_NAVY
    
    p_req = cell.add_paragraph()
    p_req.paragraph_format.space_before = Pt(2)
    p_req.paragraph_format.space_after = Pt(1)
    r_req_lbl = p_req.add_run("Require: ")
    r_req_lbl.font.name = 'Times New Roman'
    r_req_lbl.font.size = Pt(9.5)
    r_req_lbl.font.bold = True
    add_formatted_runs_to_paragraph(p_req, require_text, font_size=Pt(9.5))
    
    p_ens = cell.add_paragraph()
    p_ens.paragraph_format.space_before = Pt(1)
    p_ens.paragraph_format.space_after = Pt(4)
    # Add native border divider under Ensure
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="{HEX_DIVIDER}"/></w:pBdr>')
    p_ens._p.get_or_add_pPr().append(pBdr)
    
    r_ens_lbl = p_ens.add_run("Ensure: ")
    r_ens_lbl.font.name = 'Times New Roman'
    r_ens_lbl.font.size = Pt(9.5)
    r_ens_lbl.font.bold = True
    add_formatted_runs_to_paragraph(p_ens, ensure_text, font_size=Pt(9.5))
    
    # Steps
    for step in steps:
        p_step = cell.add_paragraph()
        p_step.paragraph_format.space_before = Pt(0)
        p_step.paragraph_format.space_after = Pt(1)
        
        # Check indentation in step text
        match = re.match(r'^(\d+:)\s*(.*)$', step)
        if match:
            num_lbl, step_body = match.groups()
            r_num = p_step.add_run(num_lbl.ljust(4))
            r_num.font.name = 'Consolas'
            r_num.font.size = Pt(9.0)
            r_num.font.bold = True
            r_num.font.color.rgb = COLOR_PRIMARY_NAVY
            
            # leading spaces indicate indent
            leading_spaces = len(step_body) - len(step_body.lstrip(' '))
            if leading_spaces > 0:
                p_step.paragraph_format.left_indent = Inches(0.15 * (leading_spaces // 2 + 1))
                step_body = step_body.lstrip(' ')
                
            add_formatted_runs_to_paragraph(p_step, step_body, font_size=Pt(9.0),
                                            font_name='Consolas')
        else:
            add_formatted_runs_to_paragraph(p_step, step, font_size=Pt(9.0),
                                            font_name='Consolas')
                                            
    # Spacer paragraph after algorithm box
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

print("Updated docx_builder_helpers loaded successfully.")
