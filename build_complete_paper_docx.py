"""
Complete generator for the publication-grade Word document (.docx)
Exact word-for-word reproduction of the paper manuscript:
"Personalized Federated Learning Mitigating Geographic Non-IID Disparity for
Vehicle Trajectory Prediction in Smart Cities"
"""

import os
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from docx_builder_helpers import (
    COLOR_PRIMARY_NAVY, COLOR_SECONDARY_SLATE, COLOR_BODY_TEXT, COLOR_MUTED_GRAY,
    HEX_NAVY, HEX_LIGHT_HEADER_BG, HEX_ALT_ROW_BG, HEX_CALLOUT_BG,
    set_cell_margins, set_cell_shading, set_booktabs_borders, set_borderless_table,
    set_table_fixed_grid, make_callout_box, set_algorithm_box_borders, add_formatted_runs_to_paragraph,
    add_equation_block, add_custom_heading, add_body_paragraph, add_figure_block,
    add_academic_table, add_algorithm_block
)

def load_all_markdown_tables(md_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
    table_patterns = re.findall(r'(\|[^\n]+\|\n\|[-:\s|]+\|\n(?:\|[^\n]+\|\n?)+)', md_text)
    tables = []
    for tp in table_patterns:
        lines = [l.strip() for l in tp.strip().splitlines() if l.strip()]
        hdr = [c.strip() for c in lines[0].split('|')[1:-1]]
        data = [ [c.strip() for c in l.split('|')[1:-1]] for l in lines[2:] ]
        tables.append((hdr, data))
    return tables

def create_full_paper_document():
    doc = docx.Document()
    
    # Configure Normal Style defaults
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_BODY_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)
    
    # Load all tables directly from markdown mirror
    md_file_path = r'c:\dev\papers\papers\JOURNAL_PAPER_MANUSCRIPT.md'
    md_tables = load_all_markdown_tables(md_file_path)
    
    # Page setup: Standard Letter with 0.75 in (54 pt) margins (IEEE Transactions standard)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    
    # Configure unlinked different first page header/footer
    section.different_first_page_header_footer = True
    
    # First page header: None (clean)
    # Subsequent pages header: Running title
    header = section.header
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run("Trivedi et al.: Personalized Federated Learning Mitigating Geographic Non-IID Disparity for Vehicle Trajectory Prediction")
    r_hdr.font.name = 'Times New Roman'
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.italic = True
    r_hdr.font.color.rgb = COLOR_MUTED_GRAY
    
    # Footer: Page numbering
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr = p_ftr.add_run("Page ")
    r_ftr.font.name = 'Times New Roman'
    r_ftr.font.size = Pt(9.0)
    r_ftr.font.color.rgb = COLOR_MUTED_GRAY
    fld_page = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p_ftr._p.append(fld_page)
    r_ftr2 = p_ftr.add_run(" of ")
    r_ftr2.font.name = 'Times New Roman'
    r_ftr2.font.size = Pt(9.0)
    r_ftr2.font.color.rgb = COLOR_MUTED_GRAY
    fld_numpages = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"/>')
    p_ftr._p.append(fld_numpages)
    
    # Also first page footer
    ftr_first = section.first_page_footer
    p_ff = ftr_first.paragraphs[0]
    p_ff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ff = p_ff.add_run("Page ")
    r_ff.font.name = 'Times New Roman'
    r_ff.font.size = Pt(9.0)
    r_ff.font.color.rgb = COLOR_MUTED_GRAY
    fld_page_1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p_ff._p.append(fld_page_1)
    r_ff2 = p_ff.add_run(" of ")
    r_ff2.font.name = 'Times New Roman'
    r_ff2.font.size = Pt(9.0)
    r_ff2.font.color.rgb = COLOR_MUTED_GRAY
    fld_numpages_1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"/>')
    p_ff._p.append(fld_numpages_1)

    # ---------------------------------------------------------
    # 1. TITLE & FRONT MATTER
    # ---------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(10)
    r_t = p_title.add_run("Personalized Federated Learning Mitigating Geographic Non-IID Disparity for Vehicle Trajectory Prediction in Smart Cities")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(17.0)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_PRIMARY_NAVY
    
    # Authors
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(0)
    p_auth.paragraph_format.space_after = Pt(3)
    
    def add_author_with_sup(p, name, sup):
        r = p.add_run(name)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.0)
        r.font.bold = True
        r.font.color.rgb = COLOR_BODY_TEXT
        rs = p.add_run(sup)
        rs.font.name = 'Times New Roman'
        rs.font.size = Pt(11.0)
        rs.font.bold = True
        rs.font.superscript = True
        rs.font.color.rgb = COLOR_PRIMARY_NAVY

    add_author_with_sup(p_auth, "Smit Bhavesh Trivedi", "1")
    p_auth.add_run(", ").font.size = Pt(11.0)
    add_author_with_sup(p_auth, "Nitya Ajit Thaker", "1")
    p_auth.add_run(", ").font.size = Pt(11.0)
    add_author_with_sup(p_auth, "Aryan Ramesh Jain", "1")
    p_auth.add_run(", and ").font.size = Pt(11.0)
    add_author_with_sup(p_auth, "Kalaavathi B.", "2")
    
    # Affiliations
    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff.paragraph_format.space_before = Pt(0)
    p_aff.paragraph_format.space_after = Pt(2)
    r_aff1_sup = p_aff.add_run("1")
    r_aff1_sup.font.name = 'Times New Roman'
    r_aff1_sup.font.size = Pt(9.5)
    r_aff1_sup.font.superscript = True
    r_aff1 = p_aff.add_run("School of Computer Science and Engineering, Vellore Institute of Technology, Vellore, India")
    r_aff1.font.name = 'Times New Roman'
    r_aff1.font.size = Pt(9.5)
    r_aff1.font.italic = True
    r_aff1.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    p_aff2 = doc.add_paragraph()
    p_aff2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff2.paragraph_format.space_before = Pt(0)
    p_aff2.paragraph_format.space_after = Pt(4)
    r_aff2_sup = p_aff2.add_run("2")
    r_aff2_sup.font.name = 'Times New Roman'
    r_aff2_sup.font.size = Pt(9.5)
    r_aff2_sup.font.superscript = True
    r_aff2 = p_aff2.add_run("Professor Grade 2, School of Computer Science and Engineering, Vellore Institute of Technology, Vellore, India")
    r_aff2.font.name = 'Times New Roman'
    r_aff2.font.size = Pt(9.5)
    r_aff2.font.italic = True
    r_aff2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Venue & Date
    p_venue = doc.add_paragraph()
    p_venue.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_venue.paragraph_format.space_before = Pt(0)
    p_venue.paragraph_format.space_after = Pt(10)
    r_v = p_venue.add_run("Target Venue: Scopus-Indexed Journal / IEEE Transactions on Intelligent Transportation Systems  |  Date: October 2026")
    r_v.font.name = 'Times New Roman'
    r_v.font.size = Pt(9.0)
    r_v.font.color.rgb = COLOR_MUTED_GRAY

    # ---------------------------------------------------------
    # 2. ABSTRACT & KEYWORDS CALLOUT BOX
    # ---------------------------------------------------------
    tbl_abstract = doc.add_table(rows=1, cols=1)
    tbl_abstract.alignment = WD_TABLE_ALIGNMENT.CENTER
    make_callout_box(tbl_abstract.rows[0].cells[0])
    set_table_fixed_grid(tbl_abstract, [7.0])
    cell_abs = tbl_abstract.rows[0].cells[0]
    
    p_abs = cell_abs.paragraphs[0]
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.line_spacing = 1.15
    p_abs.paragraph_format.space_before = Pt(0)
    p_abs.paragraph_format.space_after = Pt(6)
    
    r_abs_lbl = p_abs.add_run("Abstract—")
    r_abs_lbl.font.name = 'Times New Roman'
    r_abs_lbl.font.size = Pt(10.0)
    r_abs_lbl.font.bold = True
    r_abs_lbl.font.color.rgb = COLOR_PRIMARY_NAVY
    
    abstract_text = (
        "Multi-second vehicle trajectory prediction supports collision mitigation and signal control "
        "in Cooperative Intelligent Transportation Systems (C-ITS), but centralizing coordinates raises "
        "privacy and bandwidth concerns. Federated Learning (FL) across Roadside Units (RSUs) coordinates "
        "training without raw telemetry centralization, yet freeway, arterial, and downtown corridors exhibit "
        "geographically non-identically distributed (non-IID) kinematics. We benchmark and audit federated "
        "and personalized trajectory prediction across three functional zones of the Next Generation Simulation "
        "(NGSIM) corpus using contiguous trajectory assembly, kinematic physics baselines, matched training "
        "budgets, and five repeated seeds. Over five random seeds, Centralized training achieved a pooled Average "
        "Displacement Error (ADE) of 0.835 $\\pm$ 0.028 m, while FedAvg reached 0.887 $\\pm$ 0.048 m (12.2% normalized ADE) "
        "and FedPer reached 0.930 $\\pm$ 0.058 m, outperforming Constant Velocity (0.990 m) and Constant Acceleration "
        "(1.329 m) baselines. Dedicated local models achieved 0.880 $\\pm$ 0.048 m in their home zone, confirming local "
        "specialization remains competitive, but degraded by up to 5.04-fold out-of-domain. Gradient inversion audits "
        "on single victim windows ($|B|=1$) demonstrate that unperturbed gradients allow coordinate reconstruction "
        "(0.437 m error), whereas Gaussian noise perturbation ($\\sigma=1.0$) empirically increases reconstruction error "
        "to 5.198 m. Separately, formal multi-round batch DP-SGD accounting ($|B|=128, \\delta=10^{-5}$) bounds the worst-case "
        "privacy budget at $\\epsilon \\le 8.51$ (per-client range $2.46$--$8.51$; $\\epsilon \\le 2.80$ at $\\sigma=2.0$). "
        "Membership inference yielded near-chance detection (AUC 0.408, 51.0% accuracy). These findings establish empirical "
        "trade-offs between accuracy, edge bandwidth, and privacy accounting across heterogeneous corridors."
    )
    add_formatted_runs_to_paragraph(p_abs, abstract_text, font_size=Pt(10.0))
    
    p_kw = cell_abs.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.line_spacing = 1.15
    p_kw.paragraph_format.space_before = Pt(0)
    p_kw.paragraph_format.space_after = Pt(0)
    
    r_kw_lbl = p_kw.add_run("Keywords—")
    r_kw_lbl.font.name = 'Times New Roman'
    r_kw_lbl.font.size = Pt(9.5)
    r_kw_lbl.font.bold = True
    r_kw_lbl.font.color.rgb = COLOR_PRIMARY_NAVY
    
    keywords_text = "Cooperative intelligent transportation systems; roadside edge computing; gradient inversion attacks; differential privacy; spatio-temporal kinematics."
    r_kw_body = p_kw.add_run(keywords_text)
    r_kw_body.font.name = 'Times New Roman'
    r_kw_body.font.size = Pt(9.5)
    r_kw_body.font.italic = False
    
    # Spacing after callout box
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------
    # 3. SECTION 1. INTRODUCTION
    # ---------------------------------------------------------
    add_custom_heading(doc, "Section 1. Introduction", level=1)
    
    add_body_paragraph(
        doc,
        "Cooperative Intelligent Transportation Systems (C-ITS) and connected vehicular networks constitute "
        "critical infrastructure in contemporary smart cities, where continuous foresight of vehicular motion "
        "is indispensable for public safety and traffic efficiency [1]. Multi-second vehicle trajectory prediction "
        "enables autonomous vehicle navigation systems, connected roadside units, and traffic management centers "
        "to anticipate unsafe lane maneuvers, mitigate intersection collisions, and optimize adaptive traffic signal "
        "controllers [2], [3]. Modern deep sequence architectures, such as spatial-temporal graph networks and "
        "recurrent neural networks, excel at modeling complex multi-agent kinematics when supplied with continuous motion telemetry [4]."
    )
    
    add_body_paragraph(
        doc,
        "Conventionally, training these predictive models has relied upon centralizing dense vehicular telemetry, "
        "including geographic coordinates, heading angles, velocities, and acceleration profiles, within centralized "
        "municipal data centers [5]. However, continuous location streaming exposes commuter itineraries to privacy "
        "violations under municipal data protection statutes such as GDPR [6], [7], while consuming wireless cellular "
        "backhaul bandwidth and exposing monolithic databases to systemic breaches [8], [9], [10]. To address these "
        "centralization risks, Federated Learning (FL) has emerged as a decentralized paradigm, coordinating model "
        "training across distributed Roadside Units (RSUs) by exchanging parameterized gradient updates rather than raw "
        "vehicular coordinates [11], [12], [13], [14]."
    )
    
    add_body_paragraph(
        doc,
        "Nevertheless, deploying canonical federated learning within real-world smart cities encounters an unresolved "
        "obstacle: geographic non-identically and independently distributed (non-IID) mobility distributions across "
        "functional urban roadway sectors [15]. In metropolitan transportation networks, vehicular kinematics are "
        "strongly shaped by local road geometry, speed regulations, and traffic control topology. For instance, high-speed "
        "freeway corridors (e.g., US-101 [15]) are characterized by continuous high longitudinal velocities and gradual lane changes. "
        "In contrast, signalized arterial corridors (e.g., Lankershim Boulevard [15]) and downtown commercial street networks "
        "(e.g., Peachtree Street [15]) are dominated by frequent deceleration, stop-and-go queues, pedestrian interactions, and "
        "sharp perpendicular turns. When standard global federated aggregation algorithms enforce a single uniform parameter "
        "vector across such divergent kinematic environments, the resulting aggregation can produce client drift, gradient conflict, "
        "and elevated regional performance disparity [16]. Furthermore, while federated learning circumvents raw data transmissions, "
        "adversarial studies demonstrate that unperturbed model parameter updates remain vulnerable to gradient inversion and "
        "membership inference attacks [17], [18], [19]. The trade-offs between predictive accuracy, regional spatial fairness, "
        "wireless communication bandwidth, and privacy across heterogeneous urban edge deployments have not been comprehensively "
        "established in existing literature."
    )
    
    add_body_paragraph(
        doc,
        "The primary objective of this study is to benchmark and audit federated and personalized federated trajectory "
        "prediction across three heterogeneous NGSIM zones, comparing them with models trained on each zone alone, and reporting "
        "which commonly claimed benefits of federation the empirical findings support, with principal contributions summarized as follows:"
    )
    
    add_body_paragraph(
        doc,
        "• **Quantification of Geographic Non-IID Mobility Divergence**: A measurement of the kinematic distribution shift "
        "between freeway, arterial, and downtown zones of the Next Generation Simulation (NGSIM) corpus [15], using the mean pairwise "
        "Jensen-Shannon divergence over speed, acceleration, and longitudinal displacement ($\\overline{\\mathrm{JSD}} = 0.1088$) "
        "and a gradient cosine distance ($d_{\\mathrm{cos}} = 0.2790$, measured on one batch per client at initialization).",
        indent_level=1
    )
    
    add_body_paragraph(
        doc,
        "• **Three-Zone Multi-Seed Benchmark with Physical Baselines**: A comprehensive evaluation across five independent seeds "
        "(42, 43, 44, 45, 46) and matched 12-round/12-epoch training budgets comparing centralized, isolated, FedAvg, FedProx, "
        "pFedMe, and FedPer [20] models against Constant Velocity and Constant Acceleration physical baselines on verified "
        "contiguous trajectory segments. Global FedAvg attained an ADE of $0.887 \\pm 0.048$ m and FedPer reached $0.930 \\pm 0.058$ m "
        "($p=0.2998$), substantially improving over Constant Velocity (0.990 m) and Constant Acceleration (1.329 m), while isolated "
        "silos evaluated out-of-zone suffered marked performance degradation (up to 5.028 m ADE), supporting the utility of federated "
        "representation learning for robust cross-corridor deployment.",
        indent_level=1
    )
    
    add_body_paragraph(
        doc,
        "• **Continuous Gradient Inversion and Differential Privacy Accounting**: An empirical privacy audit adapting Deep Leakage "
        "from Gradients (DLG) [17] and loss-threshold Membership Inference Attacks (MIA) [19] to continuous trajectory regression, "
        "strictly decoupled from formal privacy accounting. On single victim windows ($|B|=1$), unprotected FedAvg exchanges leak coordinates "
        "(0.4373 m error), and withholding personalization heads alone does not prevent reconstruction (0.2960 m error), while Gaussian "
        "perturbation ($\\sigma=1.0$) degrades reconstruction beyond vehicle lane widths (5.1980 m error). Separately, multi-round "
        "batch DP-SGD with a Rényi Differential Privacy moments accountant ($|B|=128, \\delta=10^{-5}$) provides calibrated analytical "
        "bounds (worst-case $\\epsilon \\le 8.51$ at $\\sigma=1.0$, per-client range $2.46$--$8.51$; $\\epsilon \\le 2.80$ at $\\sigma=2.0$, "
        "range $0.82$--$2.80$; $\\epsilon \\le 38.74$ at $\\sigma=0.5$), while membership inference remains at near-chance discriminability "
        "(AUC 0.4078, 51.00% accuracy).",
        indent_level=1
    )
    
    add_body_paragraph(
        doc,
        "The remainder of this paper is organised as follows. Section 2 reviews related work. Section 3 describes the proposed methodology. "
        "Section 4 presents the experimental results. Section 5 discusses the findings and limitations, and Section 6 concludes."
    )

    # ---------------------------------------------------------
    # 4. SECTION 2. LITERATURE REVIEW
    # ---------------------------------------------------------
    add_custom_heading(doc, "Section 2. Literature Review", level=1)
    
    add_custom_heading(doc, "2.1 Deep Sequential Models for Vehicle Trajectory Forecasting", level=2)
    add_body_paragraph(
        doc,
        "Vehicle trajectory forecasting has evolved from classical physical kinematic models and social force formulations "
        "toward deep neural sequence architectures. Recurrent neural networks, specifically Long Short-Term Memory (LSTM) "
        "and Gated Recurrent Units (GRU), provide flexible temporal abstractions capable of modeling non-linear vehicle "
        "acceleration and velocity transitions over multi-second horizons [7], [21]. Recent investigations have incorporated "
        "Bahdanau attention mechanisms, spatial-temporal graph neural networks, and Transformer encoders to model multi-vehicle "
        "interdependencies and road geometry constraints [1], [22]. Nevertheless, the empirical efficacy of these deep neural "
        "networks remains contingent upon harvesting extensive datasets across large driving corpora. In municipal transportation "
        "networks, centralizing raw coordinate streams from thousands of connected vehicles and roadside sensors incurs communication "
        "overhead and exposes sensitive travel itineraries, motivating the need for decentralized edge learning paradigms [5], [8]."
    )
    
    add_custom_heading(doc, "2.2 Federated Learning in Cooperative Intelligent Transportation", level=2)
    add_body_paragraph(
        doc,
        "Federated learning allows distributed edge devices to cooperatively optimize machine learning models while retaining "
        "raw sensory observations on local hardware [11]. Within Cooperative Intelligent Transportation Systems, federated learning "
        "has been investigated for intersection traffic signal management, decentralized speed forecasting, and connected vehicle "
        "collision warning [3], [13]. Li *et al.* [23] designed a validity-aware semi-asynchronous federated framework to accommodate "
        "intermittent vehicular backhaul disconnections. Similarly, Reddy and Muralidhar [10] introduced zero-knowledge verifiable "
        "federated learning on permissioned blockchains to audit vehicular model parameter integrity. While these studies illustrate "
        "the communication and sovereignty advantages of federated edge computing, the majority evaluate models under artificially "
        "uniform or mildly synthetic data distributions, neglecting the structural kinematic shifts inherent to distinct real-world urban road facilities."
    )
    
    add_custom_heading(doc, "2.3 Non-IID Heterogeneity and Personalized Federated Learning", level=2)
    add_body_paragraph(
        doc,
        "Statistical data heterogeneity represents one of the primary theoretical and operational bottlenecks in federated optimization [12]. "
        "When data distributions across edge clients diverge, localized gradient updates guide client models toward contradictory regional "
        "parameter minima, causing client drift and degrading global aggregation stability [11], [12]. To mitigate client drift, "
        "Li *et al.* [12] introduced FedProx, which introduces an explicit proximal regularization penalty to bound local updates to "
        "the global consensus. However, enforcing strict model parameter consensus across fundamentally distinct operational environments "
        "can produce suboptimal predictive performance for participating clients."
    )
    add_body_paragraph(
        doc,
        "To reconcile global collaborative training with localized environmental requirements, Personalized Federated Learning (PFL) "
        "has gained significant traction. Arivazhagan *et al.* [20] proposed FedPer, which decouples deep neural networks into shared feature "
        "representation backbones and private localized output projection heads. Dinh *et al.* [24] formulated pFedMe, framing personalized "
        "federated learning as a bi-level optimization problem regulated by Moreau envelopes to balance local specialization against "
        "global consensus. Lee *et al.* [16] implemented layer-wise personalized parameter updates. Despite these theoretical advances, "
        "systematic evaluations analyzing how personalized federated architectures compress spatial performance disparity and defend against "
        "continuous coordinate leakage across diverse urban roadside environments remain absent."
    )
    
    add_custom_heading(doc, "2.4 Research Gap", level=2)
    add_body_paragraph(
        doc,
        "Although existing literature demonstrates the viability of federated computing for traffic forecasting, research gaps persist at the "
        "intersection of geographic non-IID kinematic distributions, regional spatial equity, and continuous trajectory privacy auditing. "
        "Existing studies either assume synthetic client distributions, overlook regional error disparities across high-speed and urban road "
        "corridors, or assume that gradient exchanges are inherently secure without conducting empirical gradient reconstruction audits. "
        "Table 1 summarizes representative related studies and the limitations of each, including this study."
    )
    
    # Table 1
    t1_headers, t1_rows = md_tables[0]
    t1_widths = [1.2, 1.8, 1.5, 2.5]
    add_academic_table(doc, t1_headers, t1_rows,
                       caption_text="Table 1: Comparison of Related Work and the Gap Addressed by This Study",
                       col_widths=t1_widths, font_size=Pt(8.5))

    # ---------------------------------------------------------
    # 5. SECTION 3. METHODOLOGY
    # ---------------------------------------------------------
    add_custom_heading(doc, "Section 3. Methodology", level=1)
    
    add_custom_heading(doc, "3.1 Problem Formulation", level=2)
    add_body_paragraph(
        doc,
        "We consider a smart city Cooperative Intelligent Transportation System comprising $K$ distributed Roadside Units (RSUs), "
        "denoted by $\\mathcal{C} = \\{c_1, c_2, \\dots, c_K\\}$, coordinated by a central RSU Gateway Server. Each RSU $c_k$ is situated at "
        "a distinct urban road facility (e.g., freeway corridor, signalized arterial, downtown intersection) and collects a localized dataset "
        "$\\mathcal{D}_k = \\{(X_i^{(k)}, Y_i^{(k)})\\}_{i=1}^{N_k}$, where $N_k = |\\mathcal{D}_k|$ denotes the number of trajectory records harvested by RSU $k$. "
        "The total number of trajectory samples across the metropolitan region is $N = \\sum_{k=1}^K N_k$. In this study $K = 3$, so the experiments are a three-zone transfer study rather than a large-scale federation."
    )
    
    add_body_paragraph(
        doc,
        "Let $T_{\\mathrm{obs}}$ denote the historical observation horizon and $T_{\\mathrm{pred}}$ denote the forecasted prediction horizon. "
        "Trajectory coordinates are sampled at a discrete frequency of $f_s = 2\\text{ Hz}$ (temporal interval $\\Delta t = 0.5\\text{ s}$) [15]. "
        "The observed historical sequence of vehicle $i$ at RSU $k$ over discrete timesteps $t \\in \\{-T_{\\mathrm{obs}}+1, \\dots, -1, 0\\}$ is defined as:"
    )
    add_equation_block(doc, r"X_i^{(k)} = [s_{-T_{\mathrm{obs}}+1}^{(i)}, \dots, s_{-1}^{(i)}, s_0^{(i)}]^\top \in \mathbb{R}^{T_{\mathrm{obs}} \times d_{\mathrm{in}}}", "(1)")
    
    add_body_paragraph(
        doc,
        "where each kinematic state vector $s_t^{(i)}$ is defined in a four-dimensional feature space ($d_{\\mathrm{in}} = 4$):"
    )
    add_equation_block(doc, r"s_t^{(i)} = [\Delta x_t^{(i)}, \Delta y_t^{(i)}, v_t^{(i)}, a_t^{(i)}]^\top", "(2)")
    
    add_body_paragraph(
        doc,
        "Here, $\\Delta x_t^{(i)} = x_t^{(i)} - x_0^{(i)}$ and $\\Delta y_t^{(i)} = y_t^{(i)} - y_0^{(i)}$ represent lateral (NGSIM `Local_X`) "
        "and longitudinal (`Local_Y`) coordinates expressed relative to the vehicle's position at reference observation timestep $t = 0$, "
        "providing translation invariance across map coordinate systems. The scalar quantities $v_t^{(i)}$ and $a_t^{(i)}$ denote instantaneous "
        "speed and acceleration (`v_Vel`, `v_Acc`), respectively. All four features and the targets are standardized with the mean and standard deviation "
        "of the pooled training windows of all clients; predictions are de-standardized to feet and multiplied by 0.3048 to obtain metres before any metric is computed."
    )
    
    add_body_paragraph(
        doc,
        "The ground-truth future trajectory sequence to be forecasted over future timesteps $t \\in \\{1, 2, \\dots, T_{\\mathrm{pred}}\\}$ is defined as:"
    )
    add_equation_block(doc, r"Y_i^{(k)} = [p_1^{(i)}, p_2^{(i)}, \dots, p_{T_{\mathrm{pred}}}^{(i)}]^\top \in \mathbb{R}^{T_{\mathrm{pred}} \times d_{\mathrm{out}}}", "(3)")
    
    add_body_paragraph(
        doc,
        "where $p_t^{(i)} = [\\Delta x_t^{(i)}, \\Delta y_t^{(i)}]^\\top$ represents future relative coordinate offsets ($d_{\\mathrm{out}} = 2$). "
        "In our system formulation, $T_{\\mathrm{obs}} = 6$ frames (3.0 s history) and $T_{\\mathrm{pred}} = 6$ frames (3.0 s future horizon). "
        "While widely cited NGSIM benchmarks frequently adopt an asymmetric 3 s observation / 5 s prediction horizon (evaluating 10 future steps), "
        "this study intentionally benchmarks a symmetric 3.0 s observation / 3.0 s prediction horizon (6 frames each at 2 Hz) to evaluate balanced "
        "encoder-decoder latent representations across functional corridors without temporal attenuation confounding early-round federated aggregation. "
        "Consequently, the quantitative errors reported herein should be assessed within this dedicated 6-second window formulation rather than cross-compared "
        "directly against 5-second horizon studies, with multi-second extension to 5.0 s horizons designated for future work."
    )
    
    add_body_paragraph(
        doc,
        "The global federated objective across all $K$ roadside units is to optimize a neural parameter vector $\\theta \\in \\mathbb{R}^d$ minimizing the composite empirical risk:"
    )
    add_equation_block(doc, r"\min_{\theta \in \mathbb{R}^d} \mathcal{F}(\theta) = \sum_{k=1}^K \frac{N_k}{N} \mathcal{F}_k(\theta)", "(4)")
    
    add_body_paragraph(
        doc,
        "where $\\mathcal{F}_k(\\theta)$ is the localized empirical loss computed over local dataset $\\mathcal{D}_k$:"
    )
    add_equation_block(doc, r"\mathcal{F}_k(\theta) = \frac{1}{N_k} \sum_{i=1}^{N_k} \mathcal{L}(f(X_i^{(k)}; \theta), Y_i^{(k)})", "(5)")
    
    add_body_paragraph(
        doc,
        "The sequence loss function $\\mathcal{L}(\\cdot)$ is formulated as the sequence Mean Squared Error (MSE):"
    )
    add_equation_block(doc, r"\mathcal{L}(\hat{Y}_i^{(k)}, Y_i^{(k)}) = \frac{1}{T_{\mathrm{pred}}} \sum_{t=1}^{T_{\mathrm{pred}}} \|\hat{p}_t^{(i)} - p_t^{(i)}\|_2^2", "(6)")

    # ---------------------------------------------------------
    # 3.2 DATA / MATERIALS
    # ---------------------------------------------------------
    add_custom_heading(doc, "3.2 Data / materials", level=2)
    add_body_paragraph(
        doc,
        "We utilize real-world vehicular motion telemetry from the Next Generation Simulation (NGSIM) program published by the Federal Highway Administration (FHWA) [15]. "
        "The benchmark uses a slice of 999,999 trajectory records partitioned across three geographically distinct functional road facilities acting as localized edge RSU clients. "
        "To ensure strict data reproducibility and eliminate cross-session track corruption, data preparation executes the following exact sequence of operations:"
    )
    
    prep_steps = [
        "1. **Ingestion and Slice Selection**: The combined NGSIM vehicle trajectory export is ingested in sequential chunk order, parsing columns `Vehicle_ID`, `Frame_ID`, `Global_Time`, `Local_X`, `Local_Y`, `v_Vel`, `v_Acc`, and `Location`. Numerical values are sanitized, and rows with missing or malformed fields are discarded. Exactly the first 333,333 valid records in file order are retained for each of the three locations (`us-101`, `lankershim`, and `peachtree`), yielding an exact 999,999-record slice without time-window or lane-specific filtering (the I-80 freeway dataset is excluded).",
        "2. **Chronological Sorting and Deduplication**: Within each location, duplicate records sharing identical `[Vehicle_ID, Global_Time, Frame_ID]` tuples are dropped, and records are sorted in ascending order of `[Vehicle_ID, Global_Time, Frame_ID]`.",
        "3. **Composite-Key Contiguity and Session Segmentation**: Because raw NGSIM records originate from multiple 15-minute recording periods where vehicle identifiers are recycled or tracking undergoes occlusion, contiguous trajectory segments are segmented chronologically. A new physical segment boundary is instantiated whenever $\\texttt{Vehicle\\_ID}_t \\ne \\texttt{Vehicle\\_ID}_{t-1}$, $\\texttt{Frame\\_ID}_t \\ne \\texttt{Frame\\_ID}_{t-1} + 1$, or $\\texttt{Global\\_Time}_t \\ne \\texttt{Global\\_Time}_{t-1} + 100$ ms (at native 10 Hz). Each continuous block receives a unique `segment_id`.",
        "4. **Duration Filtering and Vehicle Entity Sampling**: Segments containing fewer than 60 raw consecutive frames (6.0 s at 10 Hz) are discarded to ensure sufficient temporal support for observation and prediction horizons. Unique vehicle identifiers within valid segments are shuffled using fixed pseudo-random seed 42, retaining at most 600 representative vehicles per location to preserve domain balance.",
        "5. **Disjoint Vehicle-Level Splitting**: Unique vehicle identifiers within each location are partitioned into disjoint sets using a target 70% training, 15% validation, and 15% testing ratio under fixed seed 42. Partitioning strictly at the vehicle entity level guarantees zero vehicle overlap across splits, producing exact partition window counts (US-101: 1,351 training, 297 validation, 312 testing; Lankershim: 7,369 training, 1,660 validation, 1,680 testing; Peachtree: 11,193 training, 2,338 validation, 2,507 testing).",
        "6. **Temporal Downsampling and Window Extraction**: Surviving contiguous segments are downsampled by a factor of 5 (from 10 Hz to 2 Hz, $\\Delta t = 0.5$ s). Trajectory sequences of total length $T_{\\mathrm{obs}} + T_{\\mathrm{pred}} = 6 + 6 = 12$ steps (6.0 s) are extracted with a sliding stride of 3 steps (1.5 s). Coordinates are standardized to relative displacement $\\Delta p_t = p_t - p_0$ from the initial observation point."
    ]
    for ps in prep_steps:
        add_body_paragraph(doc, ps, indent_level=1)
        
    add_body_paragraph(
        doc,
        "The resulting client domain datasets are partitioned into disjoint 70% training, 15% validation, and 15% testing splits (yielding an exact pooled test suite of 4,499 windows evaluated across all five random seeds, as illustrated in Fig. 1a):"
    )
    
    clients_info = [
        "• **Client 0 (RSU-1: US-101 Highway)**: Multi-lane interstate freeway corridor in Los Angeles, California ($1,351$ training, $297$ validation, and $312$ test windows; mean speed $13.73 \\pm 8.49$ ft/s [$4.18 \\pm 2.59$ m/s]; mean acceleration $-0.35 \\pm 4.25$ ft/s$^2$ [$-0.11 \\pm 1.30$ m/s$^2$]; mean path length over the observation window $123.57 \\pm 79.28$ ft [$37.66 \\pm 24.16$ m]). Characterized by high longitudinal velocities and gradual lane merges.",
        "• **Client 1 (RSU-2: Lankershim Boulevard)**: Signalized commercial arterial corridor in Los Angeles, California ($7,369$ training, $1,660$ validation, and $1,680$ test windows; mean speed $5.68 \\pm 9.64$ ft/s [$1.73 \\pm 2.94$ m/s]; mean acceleration $-0.22 \\pm 3.53$ ft/s$^2$ [$-0.07 \\pm 1.08$ m/s$^2$]; mean path length over the observation window $22.60 \\pm 43.56$ ft [$6.89 \\pm 13.28$ m]). Governed by signal phases, stop-and-go queuing, and lane-selection maneuvers.",
        "• **Client 2 (RSU-3: Peachtree Street)**: Dense downtown commercial arterial grid in Atlanta, Georgia ($11,193$ training, $2,338$ validation, and $2,507$ test windows; mean speed $9.31 \\pm 14.11$ ft/s [$2.84 \\pm 4.30$ m/s]; mean acceleration $-0.09 \\pm 3.80$ ft/s$^2$ [$-0.03 \\pm 1.16$ m/s$^2$]; mean path length over the observation window $28.99 \\pm 57.29$ ft [$8.84 \\pm 17.46$ m]). Characterized by tight corner turns, low speeds, and roadside parking friction."
    ]
    for ci in clients_info:
        add_body_paragraph(doc, ci, indent_level=1)
        
    add_body_paragraph(
        doc,
        "*Pipeline Verification and Outlier Audit.* A previous preliminary version of this pipeline grouped records by `Vehicle_ID` alone without "
        "checking frame continuity across 15-minute recording periods, causing disparate vehicle tracks to be stitched into non-physical paths "
        "with displacement spikes. We resolved this defect by incorporating composite key segmentation and strict 100 ms contiguity verification "
        "prior to downsampling on the combined 999,999-record slice from `export.csv`. The corrected dataset eliminates track stitching, yielding "
        "maximum 0.5 s step displacements: US-101 (6.82 m), Lankershim (9.08 m), and Peachtree (14.00 m). An explicit audit of the 14.00 m maximum "
        "displacement on Peachtree revealed it at training window index 10712, where consecutive steps jump from $[-1.78, -16.50]$ ft to $[0.32, -62.38]$ ft "
        "($\\Delta y = 45.88$ ft in 0.5 s, equivalent to 28.0 m/s or 62.6 mph), following a preceding backward tracking oscillation of 13.6 ft. "
        "This artifact reflects known optical video tracking and occlusion jitter inherent to the raw FHWA video transcription at street intersections. "
        "Importantly, composite key contiguity verification prevents cross-session track stitching across 15-minute recording periods, but intentionally "
        "preserves within-session raw tracking jitter to avoid artificial data sanitization. Across the full Peachtree partition, 99.9% of 0.5 s "
        "displacements remain below 8.34 m (mean step displacement 1.42 m); retaining this isolated tracking artifact rather than post-hoc manual "
        "filtering preserves empirical data integrity. A documented sensitivity analysis demonstrates that retraining on the remaining 11,192 training "
        "windows under seed 42 with the test partition held fixed at 2,507 windows shifts global pooled test ADE by less than 0.001 m (from 0.880 m "
        "to 0.879 m for isolated own-zone models), confirming that benchmark conclusions are insensitive to this raw sensor tracking artifact. "
        "All benchmark experiments in Section 4 are evaluated on these clean, verified contiguous sequences."
    )
    
    # Figure 1
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig1_dataset_distribution.png",
        caption_title="Fig. 1.",
        caption_text="Empirical kinematic distribution profiles across NGSIM functional zones: (a) non-IID client data volume partitioning (train 70%, val 15%, test 15%); (b) velocity distributions reflecting distinct highway, arterial, and downtown regimes; (c) forward displacement profiles; and (d) representative continuous 2D relative trajectory paths without discontinuous jumps."
    )

    # ---------------------------------------------------------
    # 3.3 PROPOSED APPROACH
    # ---------------------------------------------------------
    add_custom_heading(doc, "3.3 Proposed Approach", level=2)
    add_body_paragraph(
        doc,
        "The empirical kinematic distribution profiles are depicted in Fig. 1. The methodology integrates deep sequence modeling, "
        "divergence quantification, decoupled parameter personalization, and adversarial gradient inversion auditing."
    )
    
    add_custom_heading(doc, "3.3.1 Deep Sequence-to-Sequence LSTM Backbone", level=3)
    add_body_paragraph(
        doc,
        "To model continuous trajectory sequences, each RSU deploys an encoder-decoder LSTM network. All reported results were "
        "obtained with this model; an attention variant present in the code base was not used in any experiment."
    )
    add_equation_block(doc, r"e_t = \mathrm{Dropout}(\mathrm{ReLU}(W_{\mathrm{in}} s_t + b_{\mathrm{in}}))", "(7)")
    add_equation_block(doc, r"h_t^{\mathrm{enc}}, c_t^{\mathrm{enc}} = \mathrm{LSTM}_{\mathrm{enc}}(e_t, (h_{t-1}^{\mathrm{enc}}, c_{t-1}^{\mathrm{enc}}))", "(8)")
    
    add_body_paragraph(
        doc,
        "where $W_{\\mathrm{in}} \\in \\mathbb{R}^{d_{\\mathrm{embed}} \\times d_{\\mathrm{in}}}$ and $b_{\\mathrm{in}} \\in \\mathbb{R}^{d_{\\mathrm{embed}}}$ "
        "denote input projection weights and bias ($d_{\\mathrm{embed}} = 64$), $\\mathrm{LSTM}_{\\mathrm{enc}}$ has two layers, $d_h = 64$ represents the "
        "hidden state dimensionality, and the dropout rate is 0.2. The decoder is a two-layer LSTM initialized with the final encoder states of both layers, "
        "$(h_0^{\\mathrm{dec}}, c_0^{\\mathrm{dec}}) = (h_{T_{\\mathrm{obs}}}^{\\mathrm{enc}}, c_{T_{\\mathrm{obs}}}^{\\mathrm{enc}})$. Its input at every future "
        "step is the final top-layer encoder state, so decoding is not autoregressive:"
    )
    add_equation_block(doc, r"h_t^{\mathrm{dec}}, c_t^{\mathrm{dec}} = \mathrm{LSTM}_{\mathrm{dec}}(h_{T_{\mathrm{obs}}}^{\mathrm{enc}}, (h_{t-1}^{\mathrm{dec}}, c_{t-1}^{\mathrm{dec}}))", "(9)")
    add_equation_block(doc, r"\hat{p}_t = W_2 \mathrm{ReLU}(W_1 h_t^{\mathrm{dec}} + b_1) + b_2", "(10)")
    
    add_body_paragraph(
        doc,
        "where $W_1 \\in \\mathbb{R}^{d_{\\mathrm{mlp}} \\times d_h}$, $b_1 \\in \\mathbb{R}^{d_{\\mathrm{mlp}}}$, $W_2 \\in \\mathbb{R}^{d_{\\mathrm{out}} \\times d_{\\mathrm{mlp}}}$, "
        "and $b_2 \\in \\mathbb{R}^{d_{\\mathrm{out}}}$ ($d_{\\mathrm{mlp}} = 32$, $d_{\\mathrm{out}} = 2$). The parameter count is as follows: input embedding "
        "$4 \\cdot 64 + 64 = 320$; each LSTM layer with 64 inputs and 64 hidden units $4(64 \\cdot 64 + 64 \\cdot 64 + 2 \\cdot 64) = 33,280$, giving $66,560$ for the encoder "
        "and $66,560$ for the decoder; output head $64 \\cdot 32 + 32 + 32 \\cdot 2 + 2 = 2,146$. The full network comprises 135,586 trainable parameters (529.63 KB in float32)."
    )
    
    add_custom_heading(doc, "3.3.2 Quantification of Geographic Non-IID Mobility Divergence", level=3)
    add_body_paragraph(
        doc,
        "To quantify kinematic non-IID shifts across RSU domains $\\mathcal{D}_i$ and $\\mathcal{D}_j$, we compute divergence metrics over the "
        "standardized speed, acceleration, and longitudinal displacement features of the training windows:"
    )
    add_body_paragraph(
        doc,
        "• **Kullback-Leibler (KL) Divergence**: Computed over $B=50$ histogram bins on a common range with smoothing factor $\\varepsilon = 10^{-10}$:",
        indent_level=1
    )
    add_equation_block(doc, r"D_{\mathrm{KL}}(P_i \parallel P_j) = \sum_{b=1}^B P_i(b) \ln \left( \frac{P_i(b) + \varepsilon}{P_j(b) + \varepsilon} \right)", "(11)")
    
    add_body_paragraph(
        doc,
        "• **Jensen-Shannon (JS) Divergence**: Symmetric and bounded metric evaluating divergence against midpoint $M = \\frac{1}{2}(P_i + P_j)$:",
        indent_level=1
    )
    add_equation_block(doc, r"\mathrm{JSD}(P_i \parallel P_j) = \frac{1}{2} D_{\mathrm{KL}}(P_i \parallel M) + \frac{1}{2} D_{\mathrm{KL}}(P_j \parallel M)", "(12)")
    
    add_body_paragraph(
        doc,
        "The aggregate regional heterogeneity score is the mean over the $F=3$ features (speed, acceleration, longitudinal displacement) and zone pairs:"
    )
    add_equation_block(doc, r"\overline{\mathrm{JSD}} = \frac{1}{F} \sum_{f=1}^F \frac{2}{K(K-1)} \sum_{1 \le i < j \le K} \mathrm{JSD}(P_i^{(f)} \parallel P_j^{(f)})", "(13)")
    
    add_body_paragraph(
        doc,
        "• **Parameter-Space Gradient Divergence**: Evaluates directional conflicts between client gradient vectors $g_i, g_j$ via cosine distance:",
        indent_level=1
    )
    add_equation_block(doc, r"d_{\mathrm{cos}}(g_i, g_j) = 1 - \frac{\langle g_i, g_j \rangle}{\|g_i\|_2 \|g_j\|_2}", "(14)")
    
    add_body_paragraph(
        doc,
        "The reported $d_{\\mathrm{cos}}$ is the mean over client pairs, computed from one batch of 64 training windows per client on a freshly initialized model. "
        "Pairwise KL divergences and Wasserstein-1 distances were also computed during the experiments but were not retained in the result files, so only $\\overline{\\mathrm{JSD}}$ and $d_{\\mathrm{cos}}$ are reported."
    )
    
    add_custom_heading(doc, "3.3.3 Personalized Federated Learning via FedPer", level=3)
    add_body_paragraph(
        doc,
        "To reduce client drift while capturing shared temporal dynamics, we implement Personalized Federated Learning via FedPer [20]. "
        "The model parameter vector is partitioned into:"
    )
    add_equation_block(doc, r"\theta = [w_{\mathrm{base}}, w_{\mathrm{head}}]", "(15)")
    
    add_body_paragraph(
        doc,
        "where $w_{\\mathrm{base}} \\in \\mathbb{R}^{d_{\\mathrm{base}}}$ comprises the shared spatial-temporal representation backbone "
        "($133,440$ parameters, $98.42\\%$ of model weights, $521.25\\text{ KiB}$), consisting of the input embedding, the encoder LSTM, "
        "and the decoder LSTM. The personalized component $w_{\\mathrm{head}} \\in \\mathbb{R}^{d_{\\mathrm{head}}}$ comprises the final "
        "regression projection layers ($2,146$ parameters, $1.58\\%$ of model weights, $8.38\\text{ KiB}$; combined parameter volume $529.63\\text{ KiB}$)."
    )
    add_body_paragraph(
        doc,
        "In each communication round, only the base parameters $w_{\\mathrm{base}}$ are uploaded to and aggregated by the gateway server:"
    )
    add_equation_block(doc, r"w_{\mathrm{base}}^{t+1} = \sum_{k=1}^K \frac{N_k}{N} w_{\mathrm{base}, k}^{t+1}", "(16)")
    
    add_body_paragraph(
        doc,
        "The personalized head $w_{\\mathrm{head}, k}$ is retained on-device at RSU $k$ and updated exclusively via local gradient descent on $\\mathcal{D}_k$. "
        "The execution flow of FedPer is detailed in Algorithm 1."
    )
    
    # Algorithm 1
    algo_steps = [
        "1: Server initializes shared base parameters w_base^0.",
        "2: for round t = 0, 1, ..., R-1 do",
        "3:   for each RSU client k in {1, ..., K} in parallel do",
        "4:     Download global base parameters: w_base,k^{t,0} <- w_base^t.",
        "5:     Retain local private head: w_head,k^{t,0} <- w_head,k^t.",
        "6:     for epoch e = 1, ..., E do",
        "7:       for mini-batch (X, Y) ~ D_k do",
        "8:         Y_pred <- f(X; [w_base,k^t, w_head,k^t])",
        "9:         L <- (1 / T_pred) * sum_{tau=1}^{T_pred} ||p_pred,tau - p_tau||^2",
        "10:        [w_base,k^t, w_head,k^t] <- [w_base,k^t, w_head,k^t] - eta * grad(L)",
        "11:      end for",
        "12:    end for",
        "13:    Upload updated base parameters w_base,k^{t+1} to Server.",
        "14:  end for",
        "15:  Server aggregates: w_base^{t+1} <- sum_{k=1}^K (N_k / N) * w_base,k^{t+1}.",
        "16: end for"
    ]
    add_algorithm_block(
        doc,
        title="Algorithm 1: Personalized Federated Learning via FedPer for RSU Networks",
        require_text="RSUs $\\mathcal{C} = \\{c_1, \\dots, c_K\\}$ with private datasets $\\mathcal{D}_k$, communication rounds $R$, local epochs $E$, learning rate $\\eta$.",
        ensure_text="Global base model $w_{\\mathrm{base}}^R$, localized client heads $\\{w_{\\mathrm{head}, k}^R\\}_{k=1}^K$.",
        steps=algo_steps
    )
    
    add_body_paragraph(
        doc,
        "*Note on Optimization.* While Algorithm 1 presents local parameter updates schematically via standard stochastic gradient "
        "descent for conceptual clarity, the operational implementation employs the Adam optimizer ($\\eta = 0.002, \\beta_1 = 0.9, \\beta_2 = 0.999$) "
        "with mini-batch size 128, weight decay $10^{-5}$, and gradient norm clipping at 5.0."
    )
    
    add_custom_heading(doc, "3.3.4 Continuous Gradient Inversion Audit and Differential Privacy Accounting", level=3)
    add_body_paragraph(
        doc,
        "To audit parameter privacy against honest-but-curious servers or backhaul eavesdroppers, we adapt Deep Leakage from Gradients (DLG) [17] "
        "and Improved DLG (iDLG) [18] to continuous multi-variate trajectory regression. We evaluate gradient inversion under worst-case privacy leakage "
        "conditions where the adversary intercepts parameter gradients computed on a single trajectory observation window ($|B|=1$) at the initial "
        "training round ($t=0$, before parameter dispersion across clients). The attack is evaluated across 10 independent victim trajectory windows "
        "sampled from Client 0's training partition. An adversary observing client parameter gradients $g_v = \\nabla_\\theta \\mathcal{L}(f(X_v; \\theta), Y_v)$ "
        "synthesizes dummy input kinematics $\\hat{X}$ and dummy trajectory targets $\\hat{Y}$ by minimizing:"
    )
    add_equation_block(doc, r"\min_{\hat{X}, \hat{Y}} \sum_{l=1}^L \|\nabla_{\theta_l} \mathcal{L}(f(\hat{X}; \theta), \hat{Y}) - g_{v, l}\|_2^2 + \lambda_{\mathrm{TV}} \mathcal{R}_{\mathrm{TV}}(\hat{X})", "(17)")
    
    add_body_paragraph(
        doc,
        "where Total Variation (TV) regularization $\\mathcal{R}_{\\mathrm{TV}}(\\hat{X}) = \\sum_{t=2}^{T_{\\mathrm{obs}}} \\|\\hat{X}_{t, 1:2} - \\hat{X}_{t-1, 1:2}\\|_2^2$ "
        "preserves kinematic motion continuity ($\\lambda_{\\mathrm{TV}} = 10^{-4}$). The attack executes over 150 iterations using the Adam optimizer ($\\eta=0.1$). "
        "For FedAvg, the adversary matches gradients across all model layers. For FedPer, the adversary at round 0 is assumed to know the initial architecture "
        "and initialization weights $w_{\\mathrm{head}}^0$; because the client withholds head gradient updates, the adversary matches gradients exclusively "
        "against the shared base backbone ($133,440$ parameters, 98.42% of model weights). We explicitly acknowledge that multi-round gradient inversion under "
        "accumulated unknown local head drift was not evaluated."
    )
    add_body_paragraph(
        doc,
        "To counter gradient inversion and bound privacy leakage, we incorporate Differentially Private Stochastic Gradient Descent (DP-SGD) [25] "
        "via per-sample gradient clipping and calibrated Gaussian perturbation. For each mini-batch $B \\subset \\mathcal{D}_k$ of size $|B|=128$, "
        "per-sample gradients $g_i = \\nabla_\\theta \\mathcal{L}(f(X_i; \\theta), Y_i)$ are clipped to a maximum Euclidean norm threshold $C$:"
    )
    add_equation_block(doc, r"\bar{g}_i = \frac{g_i}{\max\left(1, \, \frac{\|g_i\|_2}{C}\right)}, \quad \forall i \in B", "(18a)")
    
    add_body_paragraph(
        doc,
        "The privatized batch gradient update $\\tilde{g}$ transmitted by client $k$ is formed by aggregating clipped gradients and adding spherical Gaussian perturbation:"
    )
    add_equation_block(doc, r"\tilde{g} = \frac{1}{|B|} \left( \sum_{i \in B} \bar{g}_i + \mathcal{N}\left(0, \sigma^2 C^2 I\right) \right)", "(18b)")
    
    add_body_paragraph(
        doc,
        "where $C = 1.0$ is the clipping bound and $\\sigma \\in \\{0.5, 1.0, 2.0\\}$ denotes the noise multiplier."
    )
    add_body_paragraph(
        doc,
        "*Privacy Unit and Accountant Parameters.* The differential privacy guarantee applies at the level of individual 3-second trajectory "
        "observation windows (window-level DP), bounding the information leakage regarding any single continuous kinematic observation window "
        "recorded by an RSU. Because consecutive sliding windows from the same vehicle share overlapping frames, extending this guarantee to "
        "vehicle-level or trip-level DP across full commuter journeys would scale privacy loss by the number of windows via group privacy composition. "
        "Under local batch size $B=128$, local epochs $E=1$, and communication rounds $R=12$, local optimizer steps per client are $T_0 = 12 \\times 11 = 132$, "
        "$T_1 = 12 \\times 58 = 696$, and $T_2 = 12 \\times 88 = 1{,}056$ steps, retaining final partial batches without dropping records. "
        "Privacy parameters $(\\epsilon, \\delta)$ are computed using the official Opacus 1.6.0 Rényi Differential Privacy (RDP) moments accountant [25] "
        "for the subsampled Gaussian mechanism ($C=1.0, \\delta=10^{-5}$) under the standard Poisson subsampling approximation with nominal ratio $q_k = B / N_k$: "
        "Client 0 ($N_0 = 1{,}351, q_0 \\approx 0.0947$), Client 1 ($N_1 = 7{,}369, q_1 \\approx 0.0174$), and Client 2 ($N_2 = 11{,}193, q_2 \\approx 0.0114$). "
        "Because PyTorch DataLoaders retain partial final batches (e.g., 71 windows for US-101), modeling subsampling via the nominal full-batch ratio "
        "$q_k = B/N_k$ yields an analytical upper bound on privacy loss under the Poisson subsampling model. The federation-wide differential privacy bound "
        "is governed by the worst-case client (Client 0, US-101), attaining $\\epsilon \\le 8.51$ at $\\sigma = 1.0$ ($\\epsilon \\le 2.80$ at $\\sigma = 2.0$, "
        "and $\\epsilon \\le 38.74$ at $\\sigma = 0.5$). Evaluating the remaining clients yields $\\epsilon \\le 3.14$ for Client 1 ($\\epsilon \\le 1.03$ at "
        "$\\sigma=2.0$, $\\epsilon \\le 20.64$ at $\\sigma=0.5$) and $\\epsilon \\le 2.46$ for Client 2 ($\\epsilon \\le 0.82$ at $\\sigma=2.0$, $\\epsilon \\le 17.50$ "
        "at $\\sigma=0.5$). An earlier draft reported provisional client bounds of $\\epsilon \\le 5.46$ ($\\sigma=1.0$), $\\epsilon \\le 2.60$ ($\\sigma=2.0$), "
        "and $\\epsilon \\le 12.01$ ($\\sigma=0.5$) based on historical pre-filtering sample counts ($N_0=1{,}430, N_1=6{,}094, N_2=7{,}364$) and a simplified "
        "quadratic Taylor approximation $\\mathrm{RDP}(\\alpha) \\approx q^2 \\alpha / (2\\sigma^2)$; recomputing via Opacus 1.6.0 directly on the verified "
        "contiguous partitions and true step counts provides the verified theoretical privacy bounds reported here. In our gradient inversion auditing "
        "experiments, the adversary intercepts separately perturbed gradients computed on individual victim observation windows ($|B|=1$, with per-sample "
        "clipping $C=1.0$ and Gaussian noise $\\mathcal{N}(0, \\sigma^2 C^2 I)$ added to the victim gradient), directly isolating the reconstruction defense "
        "of the perturbation mechanism itself, rather than reconstructing from full 128-sample mini-batch updates which would inherently further dilute "
        "single-trajectory coordinate recovery. Importantly, the RDP bounds $(\\epsilon, \\delta)$ mathematically describe the iterative multi-round batch "
        "DP-SGD federation mechanism ($|B|=128$); they do not constitute a closed-form differential privacy guarantee for the single-window gradient inversion "
        "attack ($|B|=1$), which serves as an empirical stress test of coordinate reconstruction defense."
    )
    add_body_paragraph(
        doc,
        "In addition, a Membership Inference Attack (MIA) [19] is conducted on converged models using a loss-threshold classifier evaluated across 200 "
        "member training windows and 200 non-member test windows from Client 0. Member trajectories typically incur lower prediction loss $\\mathcal{L}$; "
        "prediction scores are defined as negative loss ($s = -\\mathcal{L}$) such that lower loss maps to higher membership likelihood. Under this standard "
        "negative-loss scoring orientation, an AUC of 0.4078 indicates that member and non-member loss distributions are substantially overlapping "
        "(corresponding to an AUC of 0.5922 under inverted positive polarity $s = +\\mathcal{L}$), yielding an optimal classification accuracy of 51.00% "
        "(near-random guessing). Because the MIA evaluation was performed on the shared base model representation of Client 0, both FedAvg and FedPer share "
        "the identical baseline representation evaluation on Client 0 (AUC 0.4078, 51.00% accuracy), accounting for their identical entries in Table 4. "
        "Note that DP-SGD perturbation was evaluated directly against gradient reconstruction (DLG), and its effect on membership inference was not separately "
        "measured; hence MIA is reported as not applicable (N/A) for DP-SGD rows in Table 6."
    )

    # ---------------------------------------------------------
    # 3.4 EXPERIMENTAL SETUP
    # ---------------------------------------------------------
    add_custom_heading(doc, "3.4 Experimental Setup", level=2)
    add_body_paragraph(
        doc,
        "All experiments were executed on a standardized computing testbed equipped with a 14-core Intel processor, 32 GB RAM, and PyTorch 2.x on Python 3.12. "
        "Optimization utilized the Adam optimizer with initial learning rate $\\eta = 0.002$, mini-batch size 128, weight decay $10^{-5}$, and gradient-norm "
        "clipping at 5.0. Federated benchmark runs executed for $R = 12$ communication rounds with 1 local epoch per round ($E = 1$), while centralized training "
        "and each isolated silo executed for 12 full epochs over their records, ensuring equal computational training budgets across all paradigms. To ensure "
        "statistical reliability and reproducibility, all benchmark models were evaluated across five repeated independent pseudo-random seeds (42, 43, 44, 45, 46), "
        "reporting multi-seed mean $\\pm$ standard deviation for ADE and two-tailed paired $t$-test significance against FedPer."
    )
    add_body_paragraph(
        doc,
        "Dataset partitioning strictly enforces vehicle-level disjoint splits: all continuous trajectory windows from a given vehicle (`Vehicle_ID`) are assigned "
        "exclusively to either the training set (70%), validation set (15%), or test set (15%) using a deterministic pseudo-random seed (seed 42), ensuring disjoint "
        "vehicle assignment between partitions. All five experimental evaluation seeds evaluate models on the exact invariant test window partitions "
        "(Client 0: 312 windows, Client 1: 1,680 windows, Client 2: 2,507 windows; 4,499 total test windows across 999,999 processed records)."
    )
    add_body_paragraph(
        doc,
        "We benchmark ten distinct architectural paradigms:"
    )
    
    paradigms = [
        "1. **Constant Velocity (CV)**: Non-parametric kinematic extrapolation projecting the last observed velocity linearly over future elapsed time:",
    ]
    add_body_paragraph(doc, paradigms[0], indent_level=1)
    add_equation_block(doc, r"\hat{p}_t = p_0 + \tau_t v_0, \quad t \in \{1, \dots, T_{\mathrm{pred}}\}", "(18c)")
    add_body_paragraph(
        doc,
        "where $\\tau_t = t \\Delta t$ denotes continuous prediction time ($\\Delta t = 0.5$ s at 2 Hz downsampling), $p_0 \\in \\mathbb{R}^2$ is the vehicle coordinate "
        "at the current observation step ($t=0$, normalized to relative origin $p_0 = \\mathbf{0}$ in the model frame), and $v_0 = (p_0 - p_{-1})/\\Delta t$ is the velocity "
        "vector estimated from the backward finite difference of the final observed coordinates.",
        indent_level=1
    )
    
    add_body_paragraph(
        doc,
        "2. **Constant Acceleration (CA)**: Non-parametric kinematic extrapolation incorporating quadratic acceleration profiles:",
        indent_level=1
    )
    add_equation_block(doc, r"\hat{p}_t = p_0 + \tau_t v_0 + \frac{1}{2} \tau_t^2 a_0, \quad t \in \{1, \dots, T_{\mathrm{pred}}\}", "(18d)")
    add_body_paragraph(
        doc,
        "where $a_0 = (v_0 - v_{-1})/\\Delta t = (p_0 - 2p_{-1} + p_{-2})/\\Delta t^2$ is the acceleration vector estimated from second-order finite differences across the final three observed coordinates.",
        indent_level=1
    )
    
    paradigms_rest = [
        "3. **Centralized (Pooled)**: Single global Seq2Seq LSTM trained on pooled multi-zone data for 12 epochs without explicit zone identifiers.",
        "4. **Centralized (Zone-Aware)**: Seq2Seq LSTM equipped with shared backbone features and zone-specific projection heads trained for 12 epochs.",
        "5. **FedAvg (Global FL)**: Canonical weighted parameter aggregation across all client RSUs [11].",
        "6. **FedProx**: Regularized global aggregation with proximal coefficient $\\mu = 0.01$ [12].",
        "7. **pFedMe**: Bi-level personalized optimization [24] with Moreau parameter $\\lambda = 1.0$, $K = 3$ inner steps, inner learning rate 0.01, and $\\beta = 0.2$.",
        "8. **FedPer**: Decoupled representation backbone with private localized heads [20].",
        "9. **Isolated Silos**: Standalone models trained strictly within individual RSU boundaries (US-101 Silo, Lankershim Silo, Peachtree Silo), evaluated both in-zone and out-of-zone.",
        "10. **Isolated (Own Zone)**: In-zone baseline where each RSU deploys its locally trained model exclusively within its native jurisdiction."
    ]
    for pr in paradigms_rest:
        add_body_paragraph(doc, pr, indent_level=1)
        
    add_body_paragraph(
        doc,
        "Hyperparameters for FedProx ($\\mu = 0.01$) and pFedMe ($\\lambda = 1.0, K = 3, \\beta = 0.2$) were adopted directly from standard literature recommendations [12], [24] "
        "without task-specific grid-search validation tuning on this three-zone dataset slice. Consequently, performance comparisons for FedProx and pFedMe are explicitly qualified "
        "as exploratory evaluations under default literature configurations, rather than definitive benchmarks of their peak capacity under exhaustive validation tuning. "
        "The observation and prediction horizons were configured to $T_{\\mathrm{obs}} = 3.0$ s and $T_{\\mathrm{pred}} = 3.0$ s (6 frames each at 2 Hz), forming a symmetric "
        "6.0 s trajectory analysis window. While standard NGSIM benchmarks frequently evaluate a 3.0 s observation / 5.0 s prediction horizon (10 prediction frames), this study "
        "intentionally evaluated matched 3.0 s horizons to evaluate symmetric encoder-decoder representations and consistent sequence dimensionalities; extending evaluation to full "
        "5.0 s horizons is discussed in Section 5."
    )
    
    add_body_paragraph(
        doc,
        "Model performance is evaluated across eight primary metrics:"
    )
    
    add_body_paragraph(doc, "• **Average Displacement Error (ADE)**: Mean Euclidean distance across all future timesteps:", indent_level=1)
    add_equation_block(doc, r"\mathrm{ADE} = \frac{1}{N T_{\mathrm{pred}}} \sum_{i=1}^N \sum_{t=1}^{T_{\mathrm{pred}}} \|\hat{p}_t^{(i)} - p_t^{(i)}\|_2", "(19)")
    add_body_paragraph(doc, "The horizon values ADE@1.0 s and ADE@2.0 s are cumulative means over steps 1–2 and 1–4, respectively.", indent_level=1)
    
    add_body_paragraph(doc, "• **Final Displacement Error (FDE)**: Euclidean distance at final horizon $T_{\\mathrm{pred}}$:", indent_level=1)
    add_equation_block(doc, r"\mathrm{FDE} = \frac{1}{N} \sum_{i=1}^N \|\hat{p}_{T_{\mathrm{pred}}}^{(i)} - p_{T_{\mathrm{pred}}}^{(i)}\|_2", "(20)")
    
    add_body_paragraph(doc, "• **Root Mean Squared Error (RMSE)**:", indent_level=1)
    add_equation_block(doc, r"\mathrm{RMSE} = \sqrt{\frac{1}{2 N T_{\mathrm{pred}}} \sum_{i=1}^N \sum_{t=1}^{T_{\mathrm{pred}}} \sum_{d=1}^2 (\hat{p}_{t, d}^{(i)} - p_{t, d}^{(i)})^2}", "(21)")
    
    add_body_paragraph(doc, "• **Mean Absolute Error (MAE)**:", indent_level=1)
    add_equation_block(doc, r"\mathrm{MAE} = \frac{1}{2 N T_{\mathrm{pred}}} \sum_{i=1}^N \sum_{t=1}^{T_{\mathrm{pred}}} \sum_{d=1}^2 |\hat{p}_{t, d}^{(i)} - p_{t, d}^{(i)}|", "(22)")
    
    add_body_paragraph(doc, "• **Mean Absolute Percentage Error (MAPE)**:", indent_level=1)
    add_equation_block(doc, r"\mathrm{MAPE} = \frac{100\%}{2 N T_{\mathrm{pred}}} \sum_{i=1}^N \sum_{t=1}^{T_{\mathrm{pred}}} \sum_{d=1}^2 \frac{|\hat{p}_{t, d}^{(i)} - p_{t, d}^{(i)}|}{\max(|p_{t, d}^{(i)}|, \varepsilon)}", "(23)")
    add_body_paragraph(doc, "where $\\varepsilon = 1.0\\text{ m}$ avoids division singularity near trajectory coordinate origins.", indent_level=1)
    
    add_body_paragraph(doc, "• **Normalized Average Displacement Error (NADE)**: Scale-invariant error normalized against mean cumulative trajectory path length:", indent_level=1)
    add_equation_block(doc, r"\mathrm{NADE} = \frac{\mathrm{ADE}}{\frac{1}{N} \sum_{i=1}^N \sum_{t=1}^{T_{\mathrm{pred}}} \|p_t^{(i)} - p_{t-1}^{(i)}\|_2} \times 100\%", "(24)")
    add_body_paragraph(doc, "Normalizing against actual cumulative path length traversed rather than net linear displacement ensures that curvature, deceleration loops, and turning maneuvers across urban intersections are accurately captured without distortion.", indent_level=1)
    
    add_body_paragraph(doc, "• **Fairness Coefficient of Variation (CoV)**: Spatial dispersion across zone errors $\\mu_{\\mathrm{zone}}$:", indent_level=1)
    add_equation_block(doc, r"\mathrm{CoV} = \frac{\sigma_{\mathrm{zone}}}{\mu_{\mathrm{zone}}} \times 100\%", "(25a)")
    
    add_body_paragraph(doc, "• **Zone Disparity Gap ($\\Delta_{\\mathrm{ADE}}$)**: Error spread between worst and best zones:", indent_level=1)
    add_equation_block(doc, r"\Delta_{\mathrm{ADE}} = \max_k \mathrm{ADE}_k - \min_k \mathrm{ADE}_k", "(25b)")

    # ---------------------------------------------------------
    # 6. SECTION 4. RESULTS
    # ---------------------------------------------------------
    add_custom_heading(doc, "Section 4. Results", level=1)
    add_body_paragraph(
        doc,
        "All experimental benchmarks report the mean $\\pm$ standard deviation across five repeated pseudo-random seeds (42, 43, 44, 45, 46), "
        "with two-tailed paired $t$-test significance relative to FedPer, conducted on the verified contiguous trajectory dataset."
    )
    
    add_custom_heading(doc, "4.1 Trajectory Prediction Benchmark", level=2)
    add_body_paragraph(
        doc,
        "Table 2 presents the multi-seed numerical evaluation across all evaluated model architectures and geographic domains. "
        "Visual trajectory rollouts are shown in Fig. 2, cross-domain evaluation in Fig. 3, communication volume in Fig. 4, and forecast horizon error growth in Fig. 5."
    )
    
    # Read Table 2 directly from JOURNAL_PAPER_MANUSCRIPT.md to ensure 100% exact numerical match
    with open(md_file_path, 'r', encoding='utf-8') as f:
        md_full = f.read()
        
    t2_headers, t2_data = md_tables[1]
    t2_match = re.search(r'\*(Table 2 Notes:.*?)\*', md_full, re.DOTALL)
    t2_notes = t2_match.group(1).strip() if t2_match else None
    
    # Exact 7.00-inch total width fitting standard 0.75-inch margins
    t2_widths = [1.15, 0.80, 0.46, 0.42, 0.45, 0.42, 0.45, 0.45, 0.45, 0.45, 0.50, 0.48, 0.52]
    add_academic_table(doc, t2_headers, t2_data,
                       caption_text="Table 2: Trajectory Forecasting Performance Benchmark Across Evaluated Paradigms and NGSIM Corridors",
                       notes_text=t2_notes, col_widths=t2_widths, font_size=Pt(7.0), group_interval=4)
                           
    add_body_paragraph(
        doc,
        "In terms of global displacement error, Centralized Pooled training with a matched budget of 12 epochs achieved a Global Pooled ADE of "
        "$0.835 \\pm 0.028$ m and FDE of 1.696 m (11.5% Normalized ADE). Centralized Zone-Aware training achieved $0.873 \\pm 0.055$ m ($p=0.2243$ vs. FedPer). "
        "In the federated regime over 12 communication rounds, canonical FedAvg achieved a Global Pooled ADE of $0.887 \\pm 0.048$ m (12.2% NADE), "
        "while FedProx attained $0.941 \\pm 0.068$ m and FedPer attained $0.930 \\pm 0.058$ m ($p=0.2998$ vs. FedPer for FedAvg). With the exception of pFedMe "
        "(which diverged to 4.612 m pooled ADE under uncalibrated default hyperparameters), evaluated deep learning approaches substantially outperformed "
        "kinematic physics baselines: Constant Velocity incurred a Global Pooled ADE of 0.990 m (13.6% NADE, $p=0.0808$ vs. FedPer), while Constant Acceleration "
        "degraded to 1.329 m (18.3% NADE, $p < 0.0001$). Notably, FedProx ($\\mu=0.01$) and pFedMe ($\\lambda=1.0, K=3, \\beta=0.2$) were evaluated under standard "
        "literature default hyperparameters without task-specific validation tuning, rendering these cross-method comparisons exploratory rather than a definitive ranking."
    )
    add_body_paragraph(
        doc,
        "When examining localized deployment at dedicated roadside units, the relevant baseline is a local model operating in its own home zone. "
        "Direct pooled evaluation across the exact test partitions (US-101: 312, Lankershim: 1,680, Peachtree: 2,507 windows; 4,499 windows total) demonstrated "
        "that isolated own-zone models achieved a pooled ADE of $0.880 \\pm 0.048$ m (0.998 m on US-101, 0.720 m on Lankershim, and 0.972 m on Peachtree). "
        "FedPer reached 1.107 m on US-101, 0.787 m on Lankershim, and 1.004 m on Peachtree. The paired $t$-test between isolated own-zone silos and FedPer yielded "
        "$p=0.2938$, indicating that the 0.050 m difference is not statistically significant. While a lack of statistical significance does not establish mathematical "
        "equivalence, it indicates that federation incurs no discernible accuracy penalty on native corridors while mitigating foreign-corridor transfer failure."
    )
    add_body_paragraph(
        doc,
        "Crucially, isolated data silos suffered marked kinematic degradation when evaluated outside their home jurisdiction (Fig. 3). When deployed on "
        "Peachtree downtown traffic, the US-101 freeway silo experienced acute error compounding to $5.028 \\pm 0.139$ m ADE (a 5.04-fold error increase over "
        "home-zone performance, $p < 0.0001$). Similarly, the Lankershim arterial silo reached 1.217 m on Peachtree. These empirical cross-domain penalties "
        "demonstrate that isolated edge sensing fails when encountering unfamiliar urban traffic topologies, supporting the utility of collaborative "
        "federated representation learning for robust cross-corridor transfer."
    )
    add_body_paragraph(
        doc,
        "Regarding wireless communication consumption, FedAvg exchanged 37.24 MiB and FedPer 36.65 MiB total volume over 12 rounds (uplink plus downlink across three clients). "
        "A one-off upload of the four normalized float32 features across 999,999 records required 15.26 MiB. This confirms that federated parameter exchange incurs comparable "
        "bandwidth to one-time feature aggregation while preserving raw telemetry decentralization on edge RSUs."
    )
    
    # Figures 2, 3, 4, 5
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig3_trajectory_predictions.png",
        caption_title="Fig. 2.",
        caption_text="Trajectory rollouts across three representative test cases: ground truth versus Centralized, FedAvg, and Isolated models, showing physically smooth continuous paths without discontinuous jumps under the corrected contiguous trajectory assembly pipeline. Axis scales differ between panels."
    )
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig4_cross_domain_generalization.png",
        caption_title="Fig. 3.",
        caption_text="Cross-zone ADE (m) of every model across evaluation zones and pooled, evaluated across 5 seeds. Boxed cells mark each isolated silo in its own native zone."
    )
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig5_communication_privacy_tradeoff.png",
        caption_title="Fig. 4.",
        caption_text="Total cumulative wireless communication volume across 12 rounds (MiB) versus global pooled test ADE (meters) across evaluated architectures, with inset zoom highlighting the 0.80--1.05 m high-accuracy cluster. Communication represents total uplink plus downlink volume (37.24 MiB for FedAvg; 36.65 MiB for FedPer; 0.00 MiB for isolated/physics baselines), compared with a computed one-off upload of four float32 features across 999,999 records (15.26 MiB)."
    )
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig6_horizon_error_growth.png",
        caption_title="Fig. 5.",
        caption_text="Cumulative pooled ADE up to 1, 2, and 3 s for each model and for the own-zone isolated baseline."
    )

    # ---------------------------------------------------------
    # 4.2 CONTROLLED HETEROGENEITY RESILIENCE
    # ---------------------------------------------------------
    add_custom_heading(doc, "4.2 Controlled Heterogeneity Resilience", level=2)
    add_body_paragraph(
        doc,
        "Table 3 and Fig. 6 report algorithmic performance under three controlled geographic non-IID training partitions: "
        "Low ($\\overline{\\mathrm{JSD}} = 0.0524$), Medium ($\\overline{\\mathrm{JSD}} = 0.1100$), and High ($\\overline{\\mathrm{JSD}} = 0.1335$). "
        "Medium corresponds to the natural three-zone partition. Low redistributes 30% of each zone's training windows into a shared pool to simulate "
        "lower cross-zone divergence. High retains the 70% of records whose kinematic feature profiles are farthest from other zone centroids, intensifying "
        "non-IID divergence. Test sets remain unchanged to ensure comparable evaluation. In this controlled study (evaluated at 3 rounds), FedAvg's "
        "mean zone ADE degraded from 1.019 m in the Low regime to 1.055 m in Medium and 1.094 m in High. FedPer attained 1.095 m (Low), 1.140 m (Medium), and 1.361 m (High)."
    )
    
    # Table 3
    t3_headers, t3_rows = md_tables[2]
    t3_widths = [2.5, 1.5, 1.5, 1.5]
    add_academic_table(doc, t3_headers, t3_rows,
                       caption_text="Table 3: Algorithmic Performance Under Resampled Geographic Non-IID Regimes (3 Rounds; Unweighted Mean Zone ADE)",
                       col_widths=t3_widths, font_size=Pt(9.0))
                       
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig7_heterogeneity_vs_personalization.png",
        caption_title="Fig. 6.",
        caption_text="Mean zone ADE of FedAvg and FedPer across controlled heterogeneity regimes (Low, Medium, High) alongside personalization delta (3 communication rounds)."
    )

    # ---------------------------------------------------------
    # 4.3 PRIVACY ATTACK AUDITING & SPATIAL FAIRNESS
    # ---------------------------------------------------------
    add_custom_heading(doc, "4.3 Privacy Attack Auditing & Spatial Fairness", level=2)
    add_body_paragraph(
        doc,
        "Table 4 and Fig. 7 report reconstruction metrics under Deep Leakage from Gradients with and without Gaussian noise added to the observed gradient. "
        "Table 5 and Fig. 8 detail spatial fairness metrics. Table 6 and Fig. 9 summarize the measured accuracy, communication, and privacy values."
    )
    
    # Table 4
    t4_headers, t4_rows = md_tables[3]
    t4_notes = (
        "Table 4: Privacy Auditing Results: Deep Leakage from Gradients (DLG) and DP-SGD Moments Accounting ($\\delta=10^{-5}$, $C=1.0$). "
        "Position reconstruction errors denote mean Euclidean coordinate displacement (in physical metres) between ground truth and inverted trajectories "
        "evaluated against separately perturbed single-window gradients ($|B|=1$, $C=1.0$). Privacy budgets ($\\epsilon$) are computed via the Opacus 1.6.0 "
        "Rényi Differential Privacy (RDP) moments accountant for the subsampled Gaussian mechanism ($C=1.0, \\delta=10^{-5}$) over local training horizons "
        "($T_0=132, T_1=696, T_2=1{,}056$ steps; $q_0=0.0947, q_1=0.0174, q_2=0.0114$). The reported worst-case federation bounds are governed by Client 0 "
        "($\\epsilon \\le 8.51$ for $\\sigma=1.0$), with per-client ranges given in parentheses. Provisional estimates from an earlier draft based on historical split counts "
        "($N_0=1{,}430, N_1=6{,}094, N_2=7{,}364$) and quadratic approximations yielded $\\epsilon \\le 12.01, 5.46, 2.60$. These RDP bounds characterize the multi-round batch "
        "DP-SGD federation mechanism ($|B|=128$); they do not constitute a closed-form differential privacy guarantee for the single-window gradient inversion attack ($|B|=1$), "
        "which serves as an empirical reconstruction benchmark. MIA AUC denotes Membership Inference Attack area under the ROC curve."
    )
    t4_widths = [1.0, 1.6, 0.6, 0.9, 0.6, 0.6, 1.7]
    add_academic_table(doc, t4_headers, t4_rows,
                       caption_text="Table 4: Privacy Auditing Results: Deep Leakage from Gradients (DLG) and DP-SGD Moments Accounting ($\\delta=10^{-5}$, $C=1.0$)",
                       notes_text=t4_notes, col_widths=t4_widths, font_size=Pt(8.0))
                       
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig8_privacy_reconstruction_vs_dp.png",
        caption_title="Fig. 7.",
        caption_text="DLG trajectory reconstruction error (metres) and cosine similarity as a function of the Gaussian noise multiplier $\\sigma$ under DP-SGD gradient perturbation ($C=1.0$). Annotations note corresponding worst-case batch DP-SGD accounting bounds ($\\delta=10^{-5}$); single-window empirical reconstruction stress tests are evaluated separately from batch accountant guarantees."
    )
    
    add_body_paragraph(
        doc,
        "The DLG gradient inversion attack was evaluated against individual victim trajectory windows ($|B|=1$) across 10 evaluation trials using "
        "the model parameter state at the initial communication round. Under unprotected FedAvg exchanges, DLG reconstructed input trajectory coordinates "
        "with an average position error of 0.4373 m and an input cosine similarity of 0.3724. Under FedPer, withholding output projection head gradients "
        "resulted in an inverted position error of 0.2960 m (cosine similarity: 0.5330), demonstrating that withholding output layers alone does not prevent "
        "trajectory reconstruction when recurrent backbone gradients remain exposed. Conversely, applying DP-SGD Gaussian noise perturbation at clipping "
        "threshold $C=1.0$ progressively degraded reconstruction fidelity: $\\sigma=0.5$ increased position error to 2.0208 m ($\\epsilon \\le 38.74$), "
        "$\\sigma=1.0$ increased error to 5.1980 m with near-zero cosine similarity (0.0508, worst-case $\\epsilon \\le 8.51$, per-client range $2.46$--$8.51$), "
        "and $\\sigma=2.0$ raised error to 9.9000 m ($\\epsilon \\le 2.80$, range $0.82$--$2.80$ at $\\delta=10^{-5}$ across local horizons), exceeding standard "
        "lane-width dimensions (3.0–3.7 m). Membership Inference Attacks on converged models, evaluated using a negative-loss score orientation ($s = -\\mathcal{L}$) "
        "across 200 member and 200 non-member test windows, yielded an AUC of 0.4078 (corresponding to 0.5922 under positive polarity) and 51.00% accuracy, confirming "
        "near-chance membership discriminability. Note that DP-SGD was evaluated specifically as a defense against gradient inversion, and its impact on membership "
        "inference was not separately measured (reported as N/A in Table 6)."
    )
    
    # Table 5
    t5_headers, t5_rows = md_tables[4]
    t5_notes = (
        "Table 5: Personalization, Spatial Disparity, and Algorithmic Fairness Analysis Across Evaluated Paradigms. "
        "Performance values averaged across 5 pseudo-random seeds. Worst, best, and mean zone metrics correspond to the three distinct evaluation jurisdictions (US-101 highway, Lankershim arterial, Peachtree downtown)."
    )
    t5_widths = [1.9, 0.85, 0.85, 0.85, 0.85, 0.85, 0.85]
    add_academic_table(doc, t5_headers, t5_rows,
                       caption_text="Table 5: Personalization, Spatial Disparity, and Algorithmic Fairness Analysis Across Evaluated Paradigms",
                       notes_text=t5_notes, col_widths=t5_widths, font_size=Pt(8.5))
                       
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig9_fairness_disparity_comparison.png",
        caption_title="Fig. 8.",
        caption_text="Spatial fairness and regional disparity across evaluated paradigms: (a) Regional performance spread (best-zone, corridor mean, and worst-zone ADE in meters); (b) Inter-zone disparity gap (worst-zone minus best-zone ADE in meters) with inset detail for clustered low-disparity regimes, including the isolated own-zone baseline in gray."
    )
    
    add_body_paragraph(
        doc,
        "In spatial fairness evaluation (Table 5 and Fig. 8), Centralized Pooled training achieved an inter-zone disparity gap of 0.253 m with a fairness CoV "
        "of 12.4%. FedAvg achieved a disparity gap of 0.293 m and CoV of 13.3% (worst-zone ADE 1.047 m on US-101, best-zone ADE 0.754 m on Lankershim). "
        "FedProx yielded a disparity gap of 0.271 m (CoV 12.1%), while FedPer achieved a disparity gap of 0.320 m (CoV 13.8%). In contrast, isolated models "
        "evaluated out-of-zone exhibited massive unfairness (US-101 silo reached CoV of 64.0% with a 4.030 m disparity gap). Isolated own-zone training achieved "
        "a balanced disparity gap of 0.278 m and CoV of 14.0%."
    )
    
    # Table 6
    t6_headers, t6_rows = md_tables[5]
    t6_notes = (
        "Table 6: Integrated Accuracy, Communication, Fairness, and Privacy Measurements. "
        "Communication denotes total cumulative transmission over 12 rounds in MiB ($2^{20}$ bytes). "
        "Centralized communication represents the computed one-off upload of only the four processed float32 model features across 999,999 records (15.26 MiB), "
        "excluding raw identifiers and preprocessing fields. DLG reconstruction evaluates inversion against separately perturbed single-window gradients "
        "($|B|=1$, $C=1.0$). Privacy budget $\\epsilon$ is evaluated using the Opacus 1.6.0 RDP moments accountant for subsampled Gaussian mechanism at "
        "$\\delta=10^{-5}$ ($C=1.0$), reporting worst-case client bound and per-client range across local training horizons ($T_0=132, T_1=696, T_2=1{,}056$). "
        "Single-window DLG reconstructions ($|B|=1$) serve as an empirical gradient leakage stress test, distinct from multi-round batch DP-SGD accounting "
        "($|B|=128$). N/A denotes configurations where metrics are not applicable or unmeasured."
    )
    t6_widths = [1.35, 0.65, 0.65, 0.65, 0.65, 0.65, 0.65, 0.55, 1.20]
    add_academic_table(doc, t6_headers, t6_rows,
                       caption_text="Table 6: Integrated Accuracy, Communication, Fairness, and Privacy Measurements",
                       notes_text=t6_notes, col_widths=t6_widths, font_size=Pt(7.5))
                       
    add_figure_block(
        doc,
        image_path=r"c:\dev\papers\papers\plots\fig10_pareto_tradeoff_frontier.png",
        caption_title="Fig. 9.",
        caption_text="Integrated empirical trade-offs: (a) Global test ADE (m) versus cumulative edge communication volume (MiB) across evaluated architectures, with inset zoom highlighting the 0.80--1.05 m high-performance frontier; (b) Empirical DLG trajectory reconstruction position error (m) across privacy mechanisms (Unprotected FedAvg, FedPer head withheld, DP-SGD $\\sigma \\in \\{0.5, 1.0, 2.0\\}$) benchmarked against arterial (3.0 m) and highway (3.7 m) lane tolerance boundaries, with separate formal batch DP-SGD moments accounting ($|B|=128, \\delta=10^{-5}$) summarized in the reference card. Centralized communication represents the computed one-off upload of the four normalized float32 features across 999,999 records (15.26 MiB)."
    )
    
    add_body_paragraph(
        doc,
        "The integrated matrix (Table 6 and Fig. 9) synthesizes accuracy, communication volume, spatial fairness, and differential privacy accounting "
        "bounds across all evaluated paradigms. Pooled Centralized training achieves the lowest global displacement error (0.835 m ADE, 12.4% CoV) under "
        "a one-off feature upload budget of 15.26 MiB. In the decentralized regime, FedAvg (0.887 m ADE, 13.3% CoV) and FedPer (0.930 m ADE, 13.8% CoV) "
        "deliver comparable predictive fidelity over 12 rounds while keeping raw trajectories local on edge RSUs. Furthermore, Gaussian gradient perturbation "
        "provides an empirical defense against single-window gradient inversion, scaling reconstruction error from 2.02 m at $\\sigma=0.5$ up to 9.90 m at "
        "$\\sigma=2.0$ (5.198 m at $\\sigma=1.0$). Separately, formal multi-round batch DP-SGD moments accounting yields worst-case client privacy bounds of "
        "$\\epsilon \\le 38.74$ ($\\sigma=0.5$), $\\epsilon \\le 8.51$ ($\\sigma=1.0$, per-client range $2.46$--$8.51$), and $\\epsilon \\le 2.80$ ($\\sigma=2.0$)."
    )

    # ---------------------------------------------------------
    # 7. SECTION 5. DISCUSSION
    # ---------------------------------------------------------
    add_custom_heading(doc, "Section 5. Discussion", level=1)
    
    add_body_paragraph(
        doc,
        "The empirical results across five independent random seeds demonstrate that collaborative federated learning closely matches the predictive accuracy "
        "of centralized pooling while preserving edge data sovereignty. Over a matched budget of 12 epochs or rounds, Centralized Pooled training achieved a "
        "Global Pooled ADE of $0.835 \\pm 0.028$ m (statistically significantly outperforming FedPer with $p=0.0325$), while canonical FedAvg attained "
        "$0.887 \\pm 0.048$ m (exhibiting no statistically significant difference from FedPer with $p=0.2998$ on pooled ADE), and FedPer reached $0.930 \\pm 0.058$ m. "
        "With the exception of pFedMe (which diverged to 4.612 m pooled ADE under uncalibrated literature default hyperparameters), evaluated deep sequential "
        "architectures substantially outperformed kinematic extrapolation baselines: Constant Velocity yielded a pooled ADE of 0.990 m (13.6% NADE, $p=0.0808$ vs. FedPer) "
        "and Constant Acceleration degraded to 1.329 m (18.3% NADE, $p < 0.0001$). Under contiguous trajectory assembly and standardized coordinates, FedAvg's "
        "shared coordinate projection benefited from aggregating temporal gradients across all three corridors, achieving slight global advantages over FedPer's "
        "locally specialized heads. Notably, FedProx ($\\mu=0.01$) and pFedMe ($\\lambda=1.0, K=3, \\beta=0.2$) were evaluated under standard literature default "
        "hyperparameters without task-specific validation tuning, rendering these cross-method comparisons exploratory rather than a definitive ranking. "
        "Crucially, both federated formulations ensured robust cross-corridor generalization, mitigating the marked performance degradation observed when isolated "
        "models operate outside their training environments."
    )
    add_body_paragraph(
        doc,
        "These findings clarify the operational roles of federated, centralized, and localized models in smart city traffic infrastructure. When evaluated "
        "within an RSU's dedicated home corridor, isolated models attained a pooled ADE of $0.880 \\pm 0.048$ m, with the paired $t$-test against FedPer (0.930 m) "
        "yielding $p=0.2938$. While failure to reject the null hypothesis does not establish mathematical equivalence, it indicates that federation incurs no "
        "discernible accuracy penalty on native corridors while safeguarding against marked cross-zone performance drops. Standalone single-zone silos exhibited "
        "acute degradation when deployed on foreign corridors; most notably, the US-101 highway silo degraded by 5.04-fold to $5.028 \\pm 0.139$ m ADE on downtown "
        "Peachtree traffic ($p < 0.0001$). Collaborative federated learning is therefore valuable not to outperform well-trained isolated models on their familiar "
        "home terrain, but to ensure operational robustness and prevent transfer failures across heterogeneous urban topologies. Furthermore, our gradient inversion "
        "audit aligns with findings by Zhu et al. [17]: raw gradient exchanges allow accurate coordinate reconstruction (0.4373 m error), and withholding output "
        "projection heads in FedPer provides negligible defense (0.2960 m error). Effective privacy defense requires calibrated gradient perturbation, where "
        "empirical noise injection ($\\sigma=1.0$) degrades single-window coordinate reconstruction to 5.1980 m error, while separate formal batch DP-SGD moments "
        "accounting bounds the worst-case client privacy budget at $\\epsilon \\le 8.51$ ($\\delta=10^{-5}$; per-client range $2.46$--$8.51$, and $\\epsilon \\le 2.80$ at $\\sigma=2.0$)."
    )
    add_body_paragraph(
        doc,
        "From a practical implementation perspective, these results provide concrete guidance for municipal edge deployments. Over 12 training rounds across "
        "three RSUs, FedAvg exchanged 37.24 MiB and FedPer exchanged 36.65 MiB total volume, while a one-off upload of the four normalized float32 features across "
        "999,999 records required 15.26 MiB (excluding raw CSV metadata, identifiers, timestamps, and preprocessing overhead). Although federated learning does "
        "not yield immediate communication reductions relative to a single raw feature upload at this data scale, its primary justification lies in continuous edge "
        "data sovereignty, regulatory compliance, and resilience against centralized data breaches. For municipal agencies concerned with eavesdropping along wireless "
        "V2X transmission links, integrating DP-SGD provides formal differential privacy bounds (worst-case $\\epsilon \\le 8.51$ at $\\sigma = 1.0$; client-level "
        "range $2.46$--$8.51$, and $\\epsilon \\le 2.80$ at $\\sigma=2.0$), while empirical gradient perturbation ($\\sigma \\ge 1.0$) degrades single-window "
        "reconstruction fidelity beyond typical vehicle lane-width dimensions (3.0--3.7 m), though downstream multi-object tracking resistance warrants task-specific empirical study."
    )
    add_body_paragraph(
        doc,
        "Several practical limitations of this study should be recognized. First, the evaluation was conducted across three distinct urban corridors representing "
        "a three-zone transfer study; expanding to dozens of federated RSU clients across partitioned road segments remains important future work. Second, we "
        "evaluated a sequence-to-sequence recurrent architecture under a 3 s observation / 3 s prediction horizon; evaluating longer multi-modal forecasting "
        "horizons (e.g., 5 s) under spatio-temporal graph attention represents a logical extension. Third, hyperparameters for FedProx ($\\mu=0.01$) and pFedMe "
        "($\\lambda=1.0, K=3, \\beta=0.2$) were adopted from standard literature defaults without client- or task-specific validation tuning, rendering these "
        "cross-method comparisons exploratory. Fourth, the differential privacy guarantee applies at the level of individual 3-second trajectory windows "
        "(window-level DP); deriving user-level or trip-level bounds across consecutive overlapping windows represents an open theoretical challenge. Finally, while "
        "the trajectory-assembly pipeline enforces frame contiguity and prevents cross-session track stitching, raw sensor tracking artifacts (such as the 14.00 m / 0.5 s "
        "displacement in Peachtree window 10712) remain preserved in the dataset to avoid artificial data sanitization. Furthermore, the NGSIM dataset is recorded under "
        "clear daylight conditions with fixed overhead cameras; performance under adverse weather, sensor noise, occlusions, and dynamic RSU disconnections requires dedicated field evaluation."
    )

    # ---------------------------------------------------------
    # 8. SECTION 6. CONCLUSION
    # ---------------------------------------------------------
    add_custom_heading(doc, "Section 6. Conclusion", level=1)
    add_body_paragraph(
        doc,
        "This study formulated, benchmarked, and audited a federated vehicle trajectory forecasting framework across smart city roadside infrastructure under "
        "geographic non-IID kinematic distributions. By enforcing contiguous trajectory assembly and conducting multi-seed evaluation across matched training budgets, "
        "we demonstrated that federated learning achieves predictive accuracy (0.887 m ADE for FedAvg, 0.930 m for FedPer) closely matching centralized training "
        "(0.835 m) and isolated own-zone baselines (0.880 m), while outperforming kinematic physics baselines (0.990 m CV, 1.329 m CA). Crucially, federated representation "
        "sharing mitigates the transfer degradation of isolated edge silos, which suffer up to a 5.04-fold error increase (5.028 m ADE) when deployed on unfamiliar "
        "corridors. Gradient inversion audits showed that withholding personalization heads alone does not prevent trajectory leakage (0.296 m error), whereas empirical "
        "Gaussian perturbation ($\\sigma=1.0$) degrades single-window coordinate reconstruction beyond standard lane widths (5.198 m error), and separate batch DP-SGD "
        "moments accounting yields calibrated privacy bounds (worst-case $\\epsilon \\le 8.51$ at $\\sigma=1.0$, per-client range $2.46$--$8.51$; $\\epsilon \\le 2.80$ at "
        "$\\sigma=2.0$, $\\delta=10^{-5}$). These findings establish quantitative empirical trade-offs between predictive accuracy, edge communication bandwidth, and "
        "differential privacy across heterogeneous smart city corridors."
    )

    # ---------------------------------------------------------
    # 9. DECLARATIONS
    # ---------------------------------------------------------
    add_custom_heading(doc, "Declarations", level=1)
    
    add_custom_heading(doc, "Acknowledgements", level=2)
    add_body_paragraph(
        doc,
        "The authors express sincere gratitude to the Federal Highway Administration (FHWA) for making the Next Generation Simulation (NGSIM) vehicular trajectory dataset publicly accessible for scientific transportation research."
    )
    
    add_custom_heading(doc, "Use of Artificial Intelligence Tools", level=2)
    add_body_paragraph(
        doc,
        "The authors used Claude (Anthropic) during manuscript preparation to assist in code auditing, reference verification against Crossref and publisher digital object identifiers, figure regeneration, and drafting text revisions in response to peer review feedback. The authors reviewed and verified all empirical computations and take full responsibility for the manuscript content."
    )
    
    add_custom_heading(doc, "Funding", level=2)
    add_body_paragraph(
        doc,
        "This research received no specific grant from any funding agency in the public, commercial, or not-for-profit sectors."
    )
    
    add_custom_heading(doc, "Conflict of Interest", level=2)
    add_body_paragraph(
        doc,
        "The authors declare no conflict of interest."
    )
    
    add_custom_heading(doc, "Data Availability", level=2)
    add_body_paragraph(
        doc,
        "The Next Generation Simulation (NGSIM) dataset is publicly available from the Federal Highway Administration (FHWA) at https://ops.fhwa.dot.gov/trafficanalysistools/ngsim.htm. The complete source code, preprocessing pipeline, model checkpoints, and benchmark evaluation scripts generated in this study are publicly available in an open-source GitHub repository at https://github.com/trivedismit13/smart_cities_project to facilitate independent reproduction."
    )

    # ---------------------------------------------------------
    # 10. REFERENCES
    # ---------------------------------------------------------
    add_custom_heading(doc, "References", level=1)
    
    references = [
        "[1] Y. Zhao, Y. Yang, H. Yu, S. Kumari, M. Amoon, S. Kumar, and Y. Ren, \"Federated learning intersection vehicle trajectory prediction scheme within digital twin,\" *IEEE Trans. Intell. Transp. Syst.*, vol. 27, no. 5, pp. 6168–6180, 2026.",
        "[2] R. Al-Huthaifi, T. Li, Z. Al-Huda, and C. Li, \"FedAGAT: Real-time traffic flow prediction based on federated community and adaptive graph attention network,\" *Inf. Sci.*, vol. 667, p. 120482, 2024.",
        "[3] C. Meese, H. Chen, W. Li, D. Lee, H. Guo, C.-C. Shen, and M. Nejad, \"Adaptive traffic prediction at the ITS edge with online models and blockchain-based federated learning,\" *IEEE Trans. Intell. Transp. Syst.*, vol. 25, no. 9, pp. 10 725–10 740, 2024.",
        "[4] X. Shen, J. Chen, S. Zhu, and R. Yan, \"A decentralized federated learning-based spatial–temporal model for freight traffic speed forecasting,\" *Expert Syst. Appl.*, vol. 238, p. 122302, 2024.",
        "[5] R. Ni, Y. Lu, B. Yang, C. Yang, and X. Liu, \"A federated pedestrian trajectory prediction model with data privacy protection,\" *Complex Intell. Syst.*, vol. 10, no. 2, pp. 1787–1799, 2024.",
        "[6] M. Han, K. Xu, S. Ma, A. Li, and H. Jiang, \"Federated learning-based trajectory prediction model with privacy preserving for intelligent vehicle,\" *Int. J. Intell. Syst.*, vol. 37, 2022.",
        "[7] A. Crivellari and Y. Shi, \"Federated LSTM-based deep learning model for privacy-preserving predictions of human trajectories across multiple data providers,\" *Ann. GIS*, 2026.",
        "[8] C. Zhang, Y. Zhu, C. Markos, S. Yu, and J. J. Q. Yu, \"A communication-efficient federated learning scheme for IoT-based traffic forecasting,\" *IEEE Internet Things J.*, vol. 9, no. 14, pp. 11 918–11 931, 2022.",
        "[9] S. Zafar, S. Jangsher, and A. Zafar, \"Federated learning-based trajectory prediction for dynamic resource allocation in moving small cell networks,\" *Veh. Commun.*, vol. 47, p. 100766, 2024.",
        "[10] K. R. Reddy and A. Muralidhar, \"Zero knowledge verifiable, semi-asynchronous federated learning for trajectory prediction on permissioned blockchain,\" *Sci. Rep.*, vol. 16, p. 18218, 2026.",
        "[11] B. McMahan, E. Moore, D. Ramage, S. Hampson, and B. A. y Arcas, \"Communication-efficient learning of deep networks from decentralized data,\" in *Proc. 20th Int. Conf. Artif. Intell. Statist. (AISTATS)*, ser. PMLR, vol. 54, 2017, pp. 1273–1282.",
        "[12] T. Li, A. K. Sahu, M. Zaheer, M. Sanjabi, A. Talwalkar, and V. Smith, \"Federated optimization in heterogeneous networks,\" *Proc. Mach. Learn. Syst. (MLSys)*, vol. 2, pp. 429–450, 2020.",
        "[13] J. Liu, L. Yang, N. Kumar, A. M. Almuhaideb, K. I. Kostromitin, and P. Zhang, \"Trajectory prediction training scheme in vehicular ad-hoc networks based on federated learning,\" *Ad Hoc Netw.*, vol. 178, p. 103917, 2025.",
        "[14] B. Li, X. Chen, T. Zhang, H. Sun, L. Dong, B. Chen, and W. Zhuang, \"Fed-SecTP: A federated-learning-based framework for secure vehicle trajectory prediction using surrounding vehicle data,\" *IEEE Internet Things J.*, 2025.",
        "[15] Federal Highway Administration (FHWA), \"Next generation simulation (NGSIM) vehicle trajectories dataset,\" U.S. Department of Transportation, Washington, DC, USA, 2006.",
        "[16] S. Lee, J. Sung, and M.-K. Shin, \"Layer-wise personalized federated learning for mobile traffic prediction,\" *IEEE Access*, vol. 12, pp. 53 126–53 140, 2024.",
        "[17] L. Zhu, Z. Liu, and S. Han, \"Deep leakage from gradients,\" in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 32, 2019, pp. 14 774–14 784.",
        "[18] B. Zhao, K. R. Mopuri, and H. Bilen, \"iDLG: Improved deep leakage from gradients,\" *arXiv:2001.02610*, 2020.",
        "[19] R. Shokri, M. Stronati, C. Song, and V. Shmatikov, \"Membership inference attacks against machine learning models,\" in *Proc. IEEE Symp. Secur. Privacy (S&P)*, 2017, pp. 3–18.",
        "[20] M. G. Arivazhagan, V. Aggarwal, A. K. Singh, and S. Choudhary, \"Federated learning with personalization layers,\" *arXiv:1912.00818*, 2019.",
        "[21] T. Mundada, S. Ramdhave, S. Jain, S. Gupta et al., \"Federated spatial-temporal traffic forecasting with VMD-enhanced graph attention and LSTM,\" *Sci. Rep.*, vol. 16, p. 8852, 2026.",
        "[22] T. Qi, L. Chen, G. Li, Y. Li, and C. Wang, \"FedAGCN: A traffic flow prediction framework based on federated learning and asynchronous graph convolutional network,\" *Appl. Soft Comput.*, vol. 138, p. 110175, 2023.",
        "[23] Y. Li, X. Xu, G. Huang, M. Yao, L. Sun, and J. Xu, \"VSFL: Trajectory prediction framework based on validity-aware semi-asynchronous federated learning in internet of vehicles,\" *Comput. Commun.*, vol. 224, pp. 106–117, 2024.",
        "[24] C. T. Dinh, N. H. Tran, and T. D. Nguyen, \"Personalized federated learning with Moreau envelopes,\" in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 33, 2020, pp. 21 394–21 405.",
        "[25] M. Abadi, A. Chu, I. Goodfellow, H. B. McMahan, I. Mironov, K. Talwar, and L. Zhang, \"Deep learning with differential privacy,\" in *Proc. ACM SIGSAC Conf. Comput. Commun. Secur. (CCS)*, 2016, pp. 308–318."
    ]
    
    for ref_str in references:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.line_spacing = 1.15
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(3)
        p_ref.paragraph_format.left_indent = Inches(0.3)
        p_ref.paragraph_format.first_line_indent = Inches(-0.3)
        add_formatted_runs_to_paragraph(p_ref, ref_str, font_size=Pt(9.0))

    # ---------------------------------------------------------
    # SAVE DESTINATIONS
    # ---------------------------------------------------------
    dest1 = r"c:\dev\papers\papers\journal_paper.docx"
    dest2 = r"c:\Users\Smit\Downloads\final_paper.docx"
    
    doc.save(dest1)
    print(f"Saved primary manuscript to: {dest1}")
    
    shutil.copy2(dest1, dest2)
    print(f"Copied finalized manuscript to: {dest2}")
    
    return dest1, dest2

if __name__ == "__main__":
    create_full_paper_document()
