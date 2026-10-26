"""
Verification script for the generated Word documents:
- c:\\dev\\papers\\papers\\journal_paper.docx
- c:\\Users\\Smit\\Downloads\\final_paper.docx
"""

import os
import sys
import zipfile
import docx

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def verify_document(filepath):
    print(f"\n=======================================================")
    print(f"VERIFYING: {filepath}")
    print(f"=======================================================")
    
    if not os.path.exists(filepath):
        print(f"ERROR: File does not exist!")
        return False
        
    size_bytes = os.path.getsize(filepath)
    print(f"File Size: {size_bytes / 1024 / 1024:.2f} MB ({size_bytes:,} bytes)")
    
    # 1. Check ZIP archive validity and media
    with zipfile.ZipFile(filepath, 'r') as zf:
        file_list = zf.namelist()
        media_files = [f for f in file_list if f.startswith('word/media/')]
        print(f"ZIP Archive valid. Embedded media files: {len(media_files)}")
        for mf in sorted(media_files):
            info = zf.getinfo(mf)
            print(f"  Media: {mf} ({info.file_size:,} bytes)")
            
    # 2. Open with python-docx
    doc = docx.Document(filepath)
    paragraphs = doc.paragraphs
    tables = doc.tables
    print(f"Total Paragraphs: {len(paragraphs)}")
    print(f"Total Tables: {len(tables)}")
    
    callout_tables = []
    eq_tables = []
    data_tables = []
    algo_tables = []
    
    for idx, t in enumerate(tables):
        r_count = len(t.rows)
        c_count = len(t.columns)
        if r_count == 1 and c_count == 1:
            text_preview = t.rows[0].cells[0].text[:40].strip()
            if "Algorithm 1" in text_preview:
                algo_tables.append((idx, r_count, c_count, text_preview))
            else:
                callout_tables.append((idx, r_count, c_count, text_preview))
        elif r_count == 1 and c_count == 2:
            eq_num = t.rows[0].cells[1].text.strip()
            eq_tables.append((idx, eq_num))
        else:
            data_tables.append((idx, r_count, c_count))
            
    print(f"\nTable Categorization:")
    print(f"  Callout Tables (Abstract): {len(callout_tables)}")
    print(f"  Algorithm Tables: {len(algo_tables)}")
    print(f"  Equation Tables: {len(eq_tables)}")
    print(f"  Data & Comparison Tables: {len(data_tables)}")
    
    print("\nData Tables Details:")
    for i, (tbl_idx, r, c) in enumerate(data_tables):
        hdr = [cell.text.strip().encode('ascii', 'replace').decode('ascii') for cell in tables[tbl_idx].rows[0].cells]
        print(f"  Table {i+1} (table #{tbl_idx}): {r} rows x {c} cols | Header: {hdr[:4]}")
        
    print(f"\nEquation Numbers Found ({len(eq_tables)} total):")
    eq_nums = [num for _, num in eq_tables]
    print("  " + ", ".join(eq_nums))
    
    # 3. Headings check
    headings = []
    for p in paragraphs:
        if p.runs and p.runs[0].bold and p.runs[0].font.size and p.runs[0].font.size.pt > 11:
            headings.append(p.text.strip().encode('ascii', 'replace').decode('ascii'))
            
    print(f"\nMajor Headings Detected ({len(headings)}):")
    for h in headings[:15]:
        print(f"  * {h}")
        
    # 4. References check
    ref_paras = [p.text.strip().encode('ascii', 'replace').decode('ascii') for p in paragraphs if p.text.strip().startswith('[') and ']' in p.text.strip()[:5]]
    print(f"\nReferences Found ({len(ref_paras)} total):")
    for rp in ref_paras[:5]:
        print(f"  {rp[:80]}...")
    print(f"  ...")
    for rp in ref_paras[-3:]:
        print(f"  {rp[:80]}...")
        
    # 5. Word count
    all_text = " ".join([p.text for p in paragraphs] + [cell.text for t in tables for row in t.rows for cell in row.cells])
    words = all_text.split()
    print(f"\nTotal Word Count in Document: {len(words):,} words")
    
    return True

if __name__ == '__main__':
    p1 = r"c:\dev\papers\papers\journal_paper.docx"
    p2 = r"c:\Users\Smit\Downloads\final_paper.docx"
    v1 = verify_document(p1)
    v2 = verify_document(p2)
    print("\nAll verifications passed!" if v1 and v2 else "\nVerification failed!")
